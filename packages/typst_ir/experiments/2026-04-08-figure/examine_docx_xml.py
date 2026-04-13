"""
Examine DOCX XML structure for figures.

Questions:
- What style names are used for figures?
- How are images embedded in paragraphs?
- What's the paragraph pattern for image + caption?
"""
from docx import Document

doc = Document("output/figure.docx")

print("=== Figure Structure ===")
print(f"Paragraphs: {len(doc.paragraphs)}")

for i, para in enumerate(doc.paragraphs):
    print(f"\n--- Paragraph {i} ---")
    print(f"Text: {para.text[:50]}...")
    print(f"Style: {para.style.name if para.style else 'None'}")
    
    # Check for image
    has_image = bool(para._element.xpath('.//a:blip'))
    print(f"Has image: {has_image}")
    
    # Check for drawing
    has_drawing = bool(para._element.xpath('.//w:drawing'))
    print(f"Has drawing: {has_drawing}")
    
    # Check for caption keywords in style
    style_lower = para.style.name.lower() if para.style else ""
    if 'caption' in style_lower or 'figure' in style_lower:
        print("  ** Figure-related style detected")
        print(f"  Style keywords: {style_lower}")

print("\n=== Pairwise Analysis ===")
for i in range(len(doc.paragraphs) - 1):
    current = doc.paragraphs[i]
    next_para = doc.paragraphs[i + 1]
    
    current_style = current.style.name.lower() if current.style else ""
    next_style = next_para.style.name.lower() if next_para.style else ""
    
    # Check for Figure + Caption pattern
    if 'figure' in current_style and 'caption' in next_style:
        print(f"Found pattern: Paragraph {i} (Figure) + Paragraph {i+1} (Caption)")
