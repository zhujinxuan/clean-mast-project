#!/usr/bin/env python3
"""Examine DOCX file for figure embedding.

This script examines the generated DOCX file to check:
1. Whether figures (images) are embedded
2. What style names are used for captions
3. How figures are structured in the document
"""

import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

# Namespaces for DOCX XML
NAMESPACES = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
}


def examine_docx(docx_path: str):
    """Examine a DOCX file for figure embedding."""
    docx_path = Path(docx_path)
    
    if not docx_path.exists():
        print(f"❌ DOCX file not found: {docx_path}")
        sys.exit(1)
    
    print("=" * 60)
    print(f"Examining DOCX: {docx_path}")
    print("=" * 60)
    
    # Open as ZIP to examine internal structure
    with zipfile.ZipFile(docx_path, 'r') as zf:
        print("\n📦 DOCX Internal Structure:")
        print("-" * 40)
        
        # List all files in the DOCX
        for name in sorted(zf.namelist()):
            if 'media' in name or 'image' in name.lower() or 'drawing' in name.lower():
                print(f"  🖼️  {name}")
            elif name.endswith('.xml'):
                if 'document' in name or 'styles' in name:
                    print(f"  📄 {name}")
    
    # Check for media folder (images)
    with zipfile.ZipFile(docx_path, 'r') as zf:
        media_files = [n for n in zf.namelist() if 'media' in n]
        if media_files:
            print(f"\n✓ Found {len(media_files)} media files (images embedded):")
            for mf in media_files:
                info = zf.getinfo(mf)
                print(f"  - {mf}: {info.file_size:,} bytes")
        else:
            print("\n❌ No media files found - images NOT embedded")
    
    # Open with python-docx for higher-level analysis
    doc = Document(str(docx_path))
    
    print("\n" + "=" * 60)
    print("Document Structure Analysis")
    print("=" * 60)
    
    # Count elements
    paragraph_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    
    print("\n📊 Element Counts:")
    print(f"  Paragraphs: {paragraph_count}")
    print(f"  Tables: {table_count}")
    
    # Analyze paragraphs for figures and captions
    print("\n📝 Paragraph Analysis:")
    print("-" * 40)
    
    caption_styles = set()
    figure_paragraphs = []
    heading_paragraphs = []
    
    for i, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else "None"
        text_preview = para.text[:50] + "..." if len(para.text) > 50 else para.text
        
        # Check for caption styles
        if 'caption' in style_name.lower():
            caption_styles.add(style_name)
            figure_paragraphs.append((i, style_name, text_preview))
        
        # Check for heading styles
        if 'heading' in style_name.lower():
            heading_paragraphs.append((i, style_name, text_preview))
        
        # Check for embedded images (drawing elements)
        drawings = para._element.findall('.//w:drawing', NAMESPACES)
        if drawings:
            print(f"  Para {i}: Has {len(drawings)} drawing(s) - Style: {style_name}")
            print(f"    Text: {text_preview}")
    
    # Print caption styles found
    if caption_styles:
        print(f"\n✓ Caption styles found: {caption_styles}")
    else:
        print("\n❌ No caption styles found")
    
    # Print figure-related paragraphs
    if figure_paragraphs:
        print(f"\n🖼️  Figure-related paragraphs ({len(figure_paragraphs)}):")
        for i, style, text in figure_paragraphs[:10]:  # Show first 10
            print(f"  Para {i}: [{style}] {text}")
    else:
        print("\n❌ No figure-related paragraphs found")
    
    # Print headings
    if heading_paragraphs:
        print(f"\n📌 Headings ({len(heading_paragraphs)}):")
        for i, style, text in heading_paragraphs[:10]:  # Show first 10
            print(f"  Para {i}: [{style}] {text}")
    
    # Check for inline shapes (another way images can be embedded)
    print("\n" + "=" * 60)
    print("Inline Shapes Analysis")
    print("=" * 60)
    
    inline_count = 0
    for para in doc.paragraphs:
        for run in para.runs:
            if run._element.findall('.//w:drawing', NAMESPACES):
                inline_count += 1
    
    print(f"Inline shapes (drawings): {inline_count}")
    
    # Check XML for drawing elements directly
    print("\n" + "=" * 60)
    print("XML Drawing Elements")
    print("=" * 60)
    
    with zipfile.ZipFile(docx_path, 'r') as zf:
        with zf.open('word/document.xml') as f:
            tree = ET.parse(f)
            root = tree.getroot()
            
            # Count drawing elements
            drawings = root.findall('.//w:drawing', NAMESPACES)
            print(f"Total <w:drawing> elements: {len(drawings)}")
            
            # Count pict elements (legacy)
            picts = root.findall('.//w:pict', NAMESPACES)
            print(f"Total <w:pict> elements: {len(picts)}")
            
            # Check for inline vs anchored
            if drawings:
                inline = root.findall('.//wp:inline', NAMESPACES)
                anchored = root.findall('.//wp:anchor', NAMESPACES)
                print(f"  Inline drawings: {len(inline)}")
                print(f"  Anchored drawings: {len(anchored)}")
    
    # Summary
    print("\n" + "=" * 60)
    print("SUMMARY")
    print("=" * 60)
    
    with zipfile.ZipFile(docx_path, 'r') as zf:
        has_media = any('media' in n for n in zf.namelist())
    
    print(f"Images embedded: {'✓ YES' if has_media else '❌ NO'}")
    print(f"Caption styles: {caption_styles if caption_styles else 'None found'}")
    print(f"Drawing elements: {len(drawings) if 'drawings' in dir() else 0}")
    
    return {
        'has_media': has_media,
        'caption_styles': list(caption_styles),
        'drawing_count': len(drawings) if 'drawings' in dir() else 0,
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        # Try default path
        docx_path = Path(__file__).parent.parent.parent.parent.parent / "test-data" / "cleanmast" / "kp-202603-report" / "rep-ts" / "report.docx"
        if not docx_path.exists():
            print("Usage: python examine_docx.py <path_to_docx>")
            print(f"       (Default path not found: {docx_path})")
            sys.exit(1)
    else:
        docx_path = sys.argv[1]
    
    result = examine_docx(docx_path)
    print(f"\nResult: {result}")