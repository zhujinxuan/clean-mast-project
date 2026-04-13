"""
Test paragraph property retrieval.

Questions:
- Can we retrieve text correctly?
- Can we compute and match content hash?
- Can we identify paragraph by position?
"""
import hashlib
import json

from docx import Document

# Load expected data
with open("input.json") as f:
    expected = json.load(f)

doc = Document("output/paragraph.docx")

# Test 1: Text retrieval
print("=== Test 1: Text Retrieval ===")
for i, (expected_para, actual_para) in enumerate(zip(expected['children'], doc.paragraphs)):
    expected_text = expected_para['text']
    actual_text = actual_para.text
    print(f"Paragraph {i}: expected='{expected_text[:30]}...' actual='{actual_text[:30]}...'")
    assert expected_text == actual_text, f"Text mismatch at paragraph {i}"

# Test 2: Content hash matching
print("\n=== Test 2: Content Hash Matching ===")
for i, (expected_para, actual_para) in enumerate(zip(expected['children'], doc.paragraphs)):
    expected_hash = hashlib.sha256(expected_para['text'].encode()).hexdigest()
    actual_hash = hashlib.sha256(actual_para.text.encode()).hexdigest()
    print(f"Paragraph {i}: expected_hash={expected_hash[:16]}... actual_hash={actual_hash[:16]}...")
    assert expected_hash == actual_hash, f"Hash mismatch at paragraph {i}"

print("\n✓ All tests passed")
