"""
Examine DOCX XML structure for paragraphs.

Questions:
- What XML elements represent paragraphs?
- What properties can we access via python-docx?
- What's the XML structure for simple text?
"""
import xml.etree.ElementTree as ET

from docx import Document

doc = Document("output/paragraph.docx")

# Examine document structure
print("=== Document Structure ===")
print(f"Paragraphs: {len(doc.paragraphs)}")

for i, para in enumerate(doc.paragraphs):
    print(f"\n--- Paragraph {i} ---")
    print(f"Text: {para.text[:50]}...")
    print(f"Style: {para.style.name if para.style else 'None'}")
    
    # Examine XML
    xml_str = ET.tostring(para._element, encoding='unicode')
    print(f"XML length: {len(xml_str)} chars")
    print(f"XML preview: {xml_str[:200]}...")
    
    # Check properties
    print(f"outline_level: {para.paragraph_format.outline_level}")
    print(f"Runs: {len(para.runs)}")
