"""
Examine DOCX XML structure for tables.

Questions:
- How are tables represented in XML?
- What's the relationship between doc.tables and doc.paragraphs?
- Can we detect table position in document order?
"""
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

doc = Document("output/table.docx")

print("=== Document Structure ===")
print(f"Paragraphs (separate): {len(doc.paragraphs)}")
print(f"Tables (separate): {len(doc.tables)}")

print("\n=== iter_inner_content Order ===")
for i, content in enumerate(doc.iter_inner_content()):
    if isinstance(content, Paragraph):
        print(f"{i}: Paragraph: {content.text[:30]}...")
    elif isinstance(content, Table):
        print(f"{i}: Table: {len(content.rows)} rows, {len(content.columns)} cols")

print("\n=== Table Cell Details ===")
for ti, table in enumerate(doc.tables):
    print(f"\nTable {ti}:")
    for ri, row in enumerate(table.rows):
        print(f"  Row {ri}:")
        for ci, cell in enumerate(row.cells):
            cell_text = " | ".join(p.text for p in cell.paragraphs)
            print(f"    Cell [{ri},{ci}]: {cell_text[:40]}...")
