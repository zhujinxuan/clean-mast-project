"""
Test equation property retrieval.

Questions:
- Can we detect display equations vs inline equations?
- What percentage of paragraphs have equations?
- Can we distinguish equation-only paragraphs?
"""
import json

from docx import Document
from docx.oxml.ns import qn

with open("input.json") as f:
    expected = json.load(f)

doc = Document("output/equation.docx")

print("=== Test 1: Detect Display Equations ===")
display_equations = []
for i, para in enumerate(doc.paragraphs):
    has_math_para = para._element.find(qn('m:oMathPara')) is not None
    if has_math_para:
        display_equations.append(i)
        print(f"Paragraph {i}: Display equation")

print(f"Total display equations: {len(display_equations)}")

print("\n=== Test 2: Detect Inline Equations ===")
inline_equations = []
for i, para in enumerate(doc.paragraphs):
    has_math = bool(para._element.xpath('.//m:oMath'))
    has_math_para = para._element.find(qn('m:oMathPara')) is not None
    
    if has_math and not has_math_para:
        inline_equations.append(i)
        print(f"Paragraph {i}: Inline equation (has text: '{para.text[:30]}...')")

print(f"Total inline equations: {len(inline_equations)}")

print("\n=== Test 3: Match Expected Equations ===")
expected_eqs = [node for node in expected['children'] if node['type'] == 'TypstEquation']
expected_display = [eq for eq in expected_eqs if eq['display']]
expected_inline = [eq for eq in expected_eqs if not eq['display']]

print(f"Expected display equations: {len(expected_display)}")
print(f"Found display equations: {len(display_equations)}")
print(f"Expected inline equations: {len(expected_inline)}")
print(f"Found inline equations: {len(inline_equations)}")

assert len(expected_display) == len(display_equations), "Display equation count mismatch"
assert len(expected_inline) == len(inline_equations), "Inline equation count mismatch"

print("\n✓ All tests passed")
