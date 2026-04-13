"""
Examine DOCX XML structure for headings.

Questions:
- Is outline_level present in XML?
- What's the relationship between style name and outline_level?
- How reliable is outline_level vs style parsing?
"""

from docx import Document
from docx.oxml.ns import qn

doc = Document("output/heading.docx")

print("=== Heading Structure ===")
print(f"Paragraphs: {len(doc.paragraphs)}")

for i, para in enumerate(doc.paragraphs):
    print(f"\n--- Paragraph {i} ---")
    print(f"Text: {para.text}")
    print(f"Style name: {para.style.name if para.style else 'None'}")
    
    # Check outline_level
    outline = para.paragraph_format.outline_level
    print(f"outline_level: {outline}")
    
    # Examine XML for outlineLvl
    pPr = para._element.find(qn('w:pPr'))
    if pPr is not None:
        outlineLvl = pPr.find(qn('w:outlineLvl'))
        if outlineLvl is not None:
            print(f"XML outlineLvl: {outlineLvl.get(qn('w:val'))}")
    
    # Check for heading style pattern
    if para.style:
        style_lower = para.style.name.lower()
        if 'heading' in style_lower:
            import re
            match = re.search(r'heading\s*(\d+)', style_lower)
            print(f"Style level: {match.group(1) if match else 'unknown'}")
