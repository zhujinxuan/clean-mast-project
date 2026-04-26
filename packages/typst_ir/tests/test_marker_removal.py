"""TDD tests for IR marker removal.

Tests remove_markers function.
"""

from __future__ import annotations

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from typst_ir.marker_detection import is_marker_paragraph
from typst_ir.marker_removal import remove_markers


def _add_marker_paragraph(doc: Document, text: str) -> None:
    """Add a marker paragraph with VerbatimChar run style."""
    try:
        verbatim = doc.styles["VerbatimChar"]
    except KeyError:
        verbatim = doc.styles.add_style("VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.style = verbatim


class TestRemoveMarkers:
    """Tests for remove_markers function."""

    def test_removes_all_marker_paragraphs(self) -> None:
        doc = Document()
        _add_marker_paragraph(doc, "irstart:block1")
        doc.add_paragraph("Content 1")
        doc.add_paragraph("Content 2")
        _add_marker_paragraph(doc, "irend:block1")
        _add_marker_paragraph(doc, "irstart:block2")
        doc.add_paragraph("Content 3")
        _add_marker_paragraph(doc, "irend:block2")

        removed = remove_markers(doc)

        assert removed == 4
        for p in doc.paragraphs:
            assert not is_marker_paragraph(p)

    def test_preserves_non_marker_content(self) -> None:
        doc = Document()
        _add_marker_paragraph(doc, "irstart:block1")
        doc.add_paragraph("Content 1")
        doc.add_paragraph("Content 2")
        _add_marker_paragraph(doc, "irend:block1")
        original_texts = [p.text for p in doc.paragraphs if not is_marker_paragraph(p)]

        remove_markers(doc)

        remaining_texts = [p.text for p in doc.paragraphs]
        assert remaining_texts == original_texts

    def test_empty_document_no_crash(self) -> None:
        doc = Document()
        removed = remove_markers(doc)
        assert removed == 0
        assert len(doc.paragraphs) == 0

    def test_mixed_markers_and_content(self) -> None:
        doc = Document()
        _add_marker_paragraph(doc, "irstart:block1")
        doc.add_paragraph("Content A")
        _add_marker_paragraph(doc, "irend:block1")
        _add_marker_paragraph(doc, "irstart:block2")
        doc.add_paragraph("Content B")
        doc.add_paragraph("Content C")
        _add_marker_paragraph(doc, "irend:block2")
        original_non_marker_texts = [
            p.text for p in doc.paragraphs if not is_marker_paragraph(p)
        ]

        removed = remove_markers(doc)

        assert removed == 4
        remaining_texts = [p.text for p in doc.paragraphs]
        assert remaining_texts == original_non_marker_texts
        assert len(remaining_texts) == 3

    def test_document_with_no_markers(self) -> None:
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        original_texts = [p.text for p in doc.paragraphs]

        removed = remove_markers(doc)

        assert removed == 0
        remaining_texts = [p.text for p in doc.paragraphs]
        assert remaining_texts == original_texts

    def test_returns_correct_count(self) -> None:
        doc = Document()
        _add_marker_paragraph(doc, "irstart:block1")
        doc.add_paragraph("Content 1")
        doc.add_paragraph("Content 2")
        doc.add_paragraph("Content 3")
        _add_marker_paragraph(doc, "irend:block1")

        removed = remove_markers(doc)

        assert removed == 2
