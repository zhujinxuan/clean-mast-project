"""Tests for synthetic DOCX fixture helpers."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from .docx_fixtures import (
    create_docx_with_markers,
    create_docx_with_sections,
)


class TestCreateDocxWithMarkers:
    """Tests for ``create_docx_with_markers``."""

    def test_creates_file_at_output_path(self, tmp_path: Path) -> None:
        output = tmp_path / "test_markers.docx"
        result = create_docx_with_markers(
            [("block1", "text", 2)], output
        )
        assert result == output
        assert output.exists()

    def test_empty_pairs_creates_minimal_docx(self, tmp_path: Path) -> None:
        output = tmp_path / "empty.docx"
        create_docx_with_markers([], output)
        doc = Document(str(output))
        assert len(doc.paragraphs) == 0

    def test_marker_count_matches_pairs(self, tmp_path: Path) -> None:
        output = tmp_path / "markers.docx"
        pairs = [
            ("block1", "text", 1),
            ("block2", "table", 2),
            ("block3", "figure", 0),
        ]
        create_docx_with_markers(pairs, output)
        doc = Document(str(output))

        expected_paragraphs = sum(2 + n for _, _, n in pairs)
        assert len(doc.paragraphs) == expected_paragraphs

    def test_start_marker_format_and_style(self, tmp_path: Path) -> None:
        output = tmp_path / "start_marker.docx"
        create_docx_with_markers([("myblock", "text", 1)], output)
        doc = Document(str(output))

        start_p = doc.paragraphs[0]
        assert start_p.text == "irstart:myblock"
        assert start_p.style.name == "FirstParagraph"
        assert all(run.style.name == "VerbatimChar" for run in start_p.runs)

    def test_end_marker_format_and_style(self, tmp_path: Path) -> None:
        output = tmp_path / "end_marker.docx"
        create_docx_with_markers([("myblock", "text", 1)], output)
        doc = Document(str(output))

        end_p = doc.paragraphs[-1]
        assert end_p.text == "irend:myblock"
        assert end_p.style.name == "Body Text"
        assert all(run.style.name == "VerbatimChar" for run in end_p.runs)

    def test_content_paragraphs_use_body_text(self, tmp_path: Path) -> None:
        output = tmp_path / "content.docx"
        create_docx_with_markers([("b1", "text", 3)], output)
        doc = Document(str(output))

        content_ps = doc.paragraphs[1:-1]
        assert len(content_ps) == 3
        for p in content_ps:
            assert p.style.name == "Body Text"
            assert "Sample text content paragraph" in p.text

    def test_multiple_blocks_are_independent(self, tmp_path: Path) -> None:
        output = tmp_path / "multi.docx"
        pairs = [("a", "text", 1), ("b", "table", 1)]
        create_docx_with_markers(pairs, output)
        doc = Document(str(output))

        texts = [p.text for p in doc.paragraphs]
        assert texts[0] == "irstart:a"
        assert texts[2] == "irend:a"
        assert texts[3] == "irstart:b"
        assert texts[5] == "irend:b"


class TestCreateDocxWithSections:
    """Tests for ``create_docx_with_sections``."""

    def test_creates_file_at_output_path(self, tmp_path: Path) -> None:
        output = tmp_path / "test_sections.docx"
        result = create_docx_with_sections(2, 1, output)
        assert result == output
        assert output.exists()

    def test_section_and_block_counts(self, tmp_path: Path) -> None:
        output = tmp_path / "counts.docx"
        num_sections = 3
        blocks_per_section = 2
        create_docx_with_sections(num_sections, blocks_per_section, output)
        doc = Document(str(output))

        expected = num_sections * (1 + blocks_per_section * 3)
        assert len(doc.paragraphs) == expected

    def test_headings_use_heading_1(self, tmp_path: Path) -> None:
        output = tmp_path / "headings.docx"
        create_docx_with_sections(2, 1, output)
        doc = Document(str(output))

        heading_ps = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert len(heading_ps) == 2
        assert heading_ps[0].text == "Section 1"
        assert heading_ps[1].text == "Section 2"

    def test_block_ids_follow_pattern(self, tmp_path: Path) -> None:
        output = tmp_path / "ids.docx"
        create_docx_with_sections(2, 2, output)
        doc = Document(str(output))

        marker_texts = [
            p.text for p in doc.paragraphs
            if p.text.startswith(("irstart:", "irend:"))
        ]
        expected_ids = [
            "irstart:section0body0", "irend:section0body0",
            "irstart:section0body1", "irend:section0body1",
            "irstart:section1body0", "irend:section1body0",
            "irstart:section1body1", "irend:section1body1",
        ]
        assert marker_texts == expected_ids

    def test_markers_have_correct_styles(self, tmp_path: Path) -> None:
        output = tmp_path / "styles.docx"
        create_docx_with_sections(1, 1, output)
        doc = Document(str(output))

        markers = [p for p in doc.paragraphs if p.text.startswith(("irstart:", "irend:"))]
        assert len(markers) == 2
        for p in markers:
            assert all(run.style.name == "VerbatimChar" for run in p.runs)

    def test_zero_sections_creates_empty_docx(self, tmp_path: Path) -> None:
        output = tmp_path / "zero.docx"
        create_docx_with_sections(0, 2, output)
        doc = Document(str(output))
        assert len(doc.paragraphs) == 0

    def test_zero_blocks_creates_headings_only(self, tmp_path: Path) -> None:
        output = tmp_path / "zero_blocks.docx"
        create_docx_with_sections(2, 0, output)
        doc = Document(str(output))
        assert len(doc.paragraphs) == 2
        assert all(p.text.startswith("Section") for p in doc.paragraphs)
