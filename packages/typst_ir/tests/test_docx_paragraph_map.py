"""Tests for docx_paragraph_map module.

Covers DocxBlockRange construction and extract_paragraph_ranges behavior.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from typst_ir.docx_paragraph_map import DocxBlockRange, extract_paragraph_ranges


def _ensure_style(doc: Document, style_name: str, style_type: WD_STYLE_TYPE) -> object:
    try:
        return doc.styles[style_name]
    except KeyError:
        return doc.styles.add_style(style_name, style_type)


class TestDocxBlockRange:
    def test_construction(self) -> None:
        r = DocxBlockRange(path="section0body2", begin=1, end=5)
        assert r.path == "section0body2"
        assert r.begin == 1
        assert r.end == 5
        assert r.table_indices == ()

    def test_with_tables(self) -> None:
        r = DocxBlockRange(
            path="section0body2table0", begin=1, end=5, table_indices=(0, 1)
        )
        assert r.table_indices == (0, 1)

    def test_is_frozen(self) -> None:
        r = DocxBlockRange(path="section0body2", begin=1, end=5)
        with pytest.raises(AttributeError):
            r.begin = 2

    def test_is_hashable(self) -> None:
        r = DocxBlockRange(path="section0body2", begin=1, end=5)
        assert hash(r) == hash(r)


class TestExtractParagraphRanges:
    def test_empty_doc_returns_empty(self, tmp_path: Path) -> None:
        doc = Document()
        path = tmp_path / "empty.docx"
        doc.save(str(path))
        result = extract_paragraph_ranges(path)
        assert result == []

    def test_single_block(self, tmp_path: Path) -> None:
        doc = Document()
        _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        verbatim = doc.styles["VerbatimChar"]

        start = doc.add_paragraph("irstart:block1")
        for run in start.runs:
            run.style = verbatim

        content = doc.add_paragraph("Content paragraph")

        end = doc.add_paragraph("irend:block1")
        for run in end.runs:
            run.style = verbatim

        path = tmp_path / "single.docx"
        doc.save(str(path))

        result = extract_paragraph_ranges(path)
        assert len(result) == 1
        assert result[0].path == "block1"
        assert result[0].begin == 1
        assert result[0].end == 2
        assert result[0].table_indices == ()

    def test_multiple_blocks(self, tmp_path: Path) -> None:
        doc = Document()
        _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        verbatim = doc.styles["VerbatimChar"]

        start1 = doc.add_paragraph("irstart:block1")
        for run in start1.runs:
            run.style = verbatim
        doc.add_paragraph("Content 1")
        end1 = doc.add_paragraph("irend:block1")
        for run in end1.runs:
            run.style = verbatim

        start2 = doc.add_paragraph("irstart:block2")
        for run in start2.runs:
            run.style = verbatim
        doc.add_paragraph("Content 2")
        end2 = doc.add_paragraph("irend:block2")
        for run in end2.runs:
            run.style = verbatim

        path = tmp_path / "multi.docx"
        doc.save(str(path))

        result = extract_paragraph_ranges(path)
        assert len(result) == 2
        assert result[0].path == "block1"
        assert result[0].begin == 1
        assert result[0].end == 2
        assert result[1].path == "block2"
        assert result[1].begin == 4
        assert result[1].end == 5

    def test_nested_blocks(self, tmp_path: Path) -> None:
        doc = Document()
        _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        verbatim = doc.styles["VerbatimChar"]

        start_outer = doc.add_paragraph("irstart:outer")
        for run in start_outer.runs:
            run.style = verbatim

        start_inner = doc.add_paragraph("irstart:inner")
        for run in start_inner.runs:
            run.style = verbatim
        doc.add_paragraph("Inner content")
        end_inner = doc.add_paragraph("irend:inner")
        for run in end_inner.runs:
            run.style = verbatim

        end_outer = doc.add_paragraph("irend:outer")
        for run in end_outer.runs:
            run.style = verbatim

        path = tmp_path / "nested.docx"
        doc.save(str(path))

        result = extract_paragraph_ranges(path)
        assert len(result) == 2
        # Inner closes first
        inner = [r for r in result if r.path == "inner"][0]
        outer = [r for r in result if r.path == "outer"][0]
        assert inner.begin == 2
        assert inner.end == 3
        assert outer.begin == 1
        assert outer.end == 4

    def test_with_table(self, tmp_path: Path) -> None:
        doc = Document()
        _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        verbatim = doc.styles["VerbatimChar"]

        start = doc.add_paragraph("irstart:table_block")
        for run in start.runs:
            run.style = verbatim

        # Add a table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Cell 1"
        table.cell(0, 1).text = "Cell 2"
        table.cell(1, 0).text = "Cell 3"
        table.cell(1, 1).text = "Cell 4"

        end = doc.add_paragraph("irend:table_block")
        for run in end.runs:
            run.style = verbatim

        path = tmp_path / "table.docx"
        doc.save(str(path))

        result = extract_paragraph_ranges(path)
        assert len(result) == 1
        assert result[0].path == "table_block"
        assert result[0].table_indices == (0,)

    def test_unmatched_start_warns(self, tmp_path: Path, caplog) -> None:
        doc = Document()
        _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        verbatim = doc.styles["VerbatimChar"]

        start = doc.add_paragraph("irstart:block1")
        for run in start.runs:
            run.style = verbatim

        path = tmp_path / "unmatched.docx"
        doc.save(str(path))

        import logging
        with caplog.at_level(logging.WARNING):
            result = extract_paragraph_ranges(path)

        assert result == []
        assert "Unmatched irstart" in caplog.text

    def test_unmatched_end_warns(self, tmp_path: Path, caplog) -> None:
        doc = Document()
        _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        verbatim = doc.styles["VerbatimChar"]

        end = doc.add_paragraph("irend:block1")
        for run in end.runs:
            run.style = verbatim

        path = tmp_path / "unmatched_end.docx"
        doc.save(str(path))

        import logging
        with caplog.at_level(logging.WARNING):
            result = extract_paragraph_ranges(path)

        assert result == []
        assert "Unmatched irend" in caplog.text
