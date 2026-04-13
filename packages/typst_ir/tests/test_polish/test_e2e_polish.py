"""End-to-end integration test for the generate_polish_ir -> apply_polish_ir pipeline.

Exercises the full pipeline with both styles and renumbering, verifying
dot-separated numbering and style application.
"""

from __future__ import annotations

from pathlib import Path

from click.testing import CliRunner
from docx import Document
from typst_ir import (
    TypstDocument,
    TypstFigure,
    TypstHeading,
    TypstParagraph,
    TypstSection,
    TypstTable,
)
from typst_ir.mapping.cli import main


def test_e2e_polish_pipeline(tmp_path: Path) -> None:
    docx_path = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("Intro text", style="Normal")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Body text", style="Normal")
    doc.add_heading("Background", level=2)
    doc.add_paragraph("Figure caption", style="Caption")
    doc.add_paragraph("Table caption", style="Caption")
    doc.save(str(docx_path))

    ir = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Introduction"),
                body=[
                    TypstParagraph(text="Body text"),
                ],
                subsections=[
                    TypstSection(
                        heading=TypstHeading(text="Background"),
                        body=[
                            TypstFigure(content="img.png", caption="Figure caption"),
                            TypstTable(
                                headers=["A"],
                                data=[["1"]],
                                caption="Table caption",
                            ),
                        ],
                        subsections=[],
                    ),
                ],
            ),
        ],
    )
    ir_json_path = tmp_path / "ir.json"
    ir_json_path.write_text(ir.model_dump_json(), encoding="utf-8")

    config_path = tmp_path / "config.toml"
    config_path.write_text(
        '\n'
        '[styles."Normal"]\n'
        'font_size = 10.0\n'
        'line_spacing = 1.15\n'
        '\n'
        '[styles."Caption"]\n'
        'font_size = 9.0\n'
        '\n'
        '[renumber]\n'
        'headings = true\n'
        'figures = true\n'
        'tables = true\n',
        encoding="utf-8",
    )

    polish_ir_path = tmp_path / "polish.json"
    output_docx_path = tmp_path / "output.docx"

    runner = CliRunner()

    result_generate = runner.invoke(
        main,
        [
            "generate-polish-ir",
            "--docx", str(docx_path),
            "--config", str(config_path),
            "--output-ir", str(polish_ir_path),
        ],
    )
    assert result_generate.exit_code == 0, result_generate.output
    assert polish_ir_path.exists()

    result_apply = runner.invoke(
        main,
        [
            "apply-polish-ir",
            "--docx", str(docx_path),
            "--polish-ir", str(polish_ir_path),
            "--ir-json", str(ir_json_path),
            "--output-docx", str(output_docx_path),
        ],
    )
    assert result_apply.exit_code == 0, result_apply.output
    assert output_docx_path.exists()

    out_doc = Document(str(output_docx_path))
    paragraph_texts = [p.text for p in out_doc.paragraphs]

    assert any(
        text.startswith("Figure ") and ":" in text and "." in text.split(":")[0]
        for text in paragraph_texts
    ), f"Expected dot-separated figure numbering in {paragraph_texts}"
    assert any(
        text.startswith("Table ") and ":" in text and "." in text.split(":")[0]
        for text in paragraph_texts
    ), f"Expected dot-separated table numbering in {paragraph_texts}"
    assert any(
        text.split()[0].rstrip(".").isdigit() if text else False
        for text in paragraph_texts
    ), f"Expected numbered heading in {paragraph_texts}"

    normal = out_doc.styles["Normal"]
    assert normal.font.size is not None
    assert normal.font.size.pt == 10.0

    caption = out_doc.styles["Caption"]
    assert caption.font.size is not None
    assert caption.font.size.pt == 9.0
