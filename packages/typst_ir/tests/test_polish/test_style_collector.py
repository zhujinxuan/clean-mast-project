"""Tests for style_collector module.

Verifies style extraction from DOCX files using python-docx fixtures.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from typst_ir.polish.style_collector import (
    StyleInfo,
    _derive_font_style,
    _extract_style_info,
    collect_styles,
    collect_styles_for_names,
)


def test_derive_font_style_all_combinations() -> None:
    """_derive_font_style returns correct style for all bold/italic combos."""
    assert _derive_font_style(True, True) == "bold_italic"
    assert _derive_font_style(True, False) == "bold"
    assert _derive_font_style(True, None) == "bold"
    assert _derive_font_style(False, True) == "italic"
    assert _derive_font_style(None, True) == "italic"
    assert _derive_font_style(False, False) == "normal"
    assert _derive_font_style(None, None) is None
    assert _derive_font_style(False, None) == "normal"
    assert _derive_font_style(None, False) == "normal"


def test_style_info_defaults() -> None:
    """StyleInfo can be created with all fields defaulting to None."""
    info = StyleInfo()
    assert info.font_name is None
    assert info.font_size is None
    assert info.font_style is None
    assert info.font_color is None
    assert info.line_spacing_before is None
    assert info.line_spacing_after is None


def test_style_info_creation_full() -> None:
    """StyleInfo can be created with all fields populated."""
    info = StyleInfo(
        font_name="Arial",
        font_size=12.0,
        font_style="bold_italic",
        font_color="#FF0000",
        line_spacing_before=6.0,
        line_spacing_after=6.0,
    )
    assert info.font_name == "Arial"
    assert info.font_size == 12.0
    assert info.font_style == "bold_italic"
    assert info.font_color == "#FF0000"
    assert info.line_spacing_before == 6.0
    assert info.line_spacing_after == 6.0


class TestCollectStyles:
    @pytest.fixture
    def styled_doc(self, tmp_path: Path) -> Path:
        pytest.importorskip("docx")
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc_path = tmp_path / "styled.docx"
        doc = Document()

        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.font.bold = True
        normal.font.italic = True
        normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        normal.paragraph_format.space_before = Pt(6)
        normal.paragraph_format.space_after = Pt(6)

        heading1 = doc.styles["Heading 1"]
        heading1.font.name = "Arial"
        heading1.font.size = Pt(16)
        heading1.font.bold = True
        heading1.font.italic = False
        heading1.paragraph_format.space_before = Pt(12)
        heading1.paragraph_format.space_after = Pt(12)

        doc.add_paragraph("Body text", style="Normal")
        doc.add_heading("Heading", level=1)
        doc.save(str(doc_path))
        return doc_path

    def test_collect_styles_extracts_all(self, styled_doc: Path) -> None:
        """collect_styles returns a dict with extracted style properties."""
        result = collect_styles(styled_doc)

        assert "Normal" in result
        normal = result["Normal"]
        assert normal.font_name == "Times New Roman"
        assert normal.font_size == 12.0
        assert normal.font_style == "bold_italic"
        assert normal.font_color == "000000"
        assert normal.line_spacing_before == 6.0
        assert normal.line_spacing_after == 6.0

        assert "Heading 1" in result
        heading1 = result["Heading 1"]
        assert heading1.font_name == "Arial"
        assert heading1.font_size == 16.0
        assert heading1.font_style == "bold"
        assert heading1.line_spacing_before == 12.0
        assert heading1.line_spacing_after == 12.0

    def test_collect_styles_returns_many(self, styled_doc: Path) -> None:
        """collect_styles includes more than just the modified styles."""
        result = collect_styles(styled_doc)
        assert len(result) > 2


class TestCollectStylesForNames:
    @pytest.fixture
    def styled_doc(self, tmp_path: Path) -> Path:
        pytest.importorskip("docx")
        from docx import Document
        from docx.shared import Pt

        doc_path = tmp_path / "styled.docx"
        doc = Document()

        normal = doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)

        heading1 = doc.styles["Heading 1"]
        heading1.font.name = "Arial"
        heading1.font.size = Pt(16)

        doc.add_paragraph("Body text", style="Normal")
        doc.add_heading("Heading", level=1)
        doc.save(str(doc_path))
        return doc_path

    def test_filter_returns_requested_styles(self, styled_doc: Path) -> None:
        """collect_styles_for_names returns only requested styles."""
        result = collect_styles_for_names(styled_doc, ["Normal", "Heading 1"])
        assert set(result.keys()) == {"Normal", "Heading 1"}
        assert result["Normal"].font_name == "Times New Roman"
        assert result["Heading 1"].font_name == "Arial"

    def test_filter_omits_missing_and_warns(self, styled_doc: Path, caplog) -> None:
        """collect_styles_for_names warns when a requested style is missing."""
        with caplog.at_level("WARNING"):
            result = collect_styles_for_names(
                styled_doc, ["Normal", "Nonexistent Style"]
            )

        assert set(result.keys()) == {"Normal"}
        assert "Nonexistent Style" in caplog.text
        assert "not found" in caplog.text


class TestExtractStyleInfoEdgeCases:
    @pytest.fixture
    def minimal_doc(self, tmp_path: Path) -> Path:
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "minimal.docx"
        doc = Document()
        doc.add_paragraph("Hello")
        doc.save(str(doc_path))
        return doc_path

    def test_extract_handles_none_properties(self, minimal_doc: Path) -> None:
        """_extract_style_info handles styles with mostly None properties."""
        pytest.importorskip("docx")
        from docx import Document

        doc = Document(str(minimal_doc))
        style = doc.styles["Normal"]
        info = _extract_style_info(style)

        assert isinstance(info, StyleInfo)
        # Some of these may have values from the built-in style defaults,
        # but the extraction should not crash on None attributes.

    def test_character_style_no_paragraph_format(self, tmp_path: Path) -> None:
        """_extract_style_info handles character styles without paragraph_format."""
        pytest.importorskip("docx")
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt

        doc_path = tmp_path / "char_style.docx"
        doc = Document()
        char_style = doc.styles.add_style("CustomChar", WD_STYLE_TYPE.CHARACTER)
        char_style.font.name = "Courier New"
        char_style.font.size = Pt(10)
        char_style.font.bold = False
        char_style.font.italic = True

        para = doc.add_paragraph("Text")
        run = para.add_run(" styled")
        run.style = char_style
        doc.save(str(doc_path))

        result = collect_styles(doc_path)
        info = result["CustomChar"]
        assert info.font_name == "Courier New"
        assert info.font_size == 10.0
        assert info.font_style == "italic"
        assert info.line_spacing_before is None
        assert info.line_spacing_after is None
