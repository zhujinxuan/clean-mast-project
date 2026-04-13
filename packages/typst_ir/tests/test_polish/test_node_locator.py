import hashlib
from pathlib import Path

import pytest
from typst_ir.mapping.types import DocxTypstMapping, MappingEntry, NodeID


def _create_node_id(path: str, doc_content: str) -> NodeID:
    full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]
    node_hash = hashlib.sha256(path.encode()).hexdigest()[:16]
    return NodeID(path=path, full_hash=full_hash, node_hash=node_hash)


class TestLocateNodes:
    @pytest.fixture
    def multi_para_doc(self, tmp_path: Path) -> Path:
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "multi.docx"
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        doc.add_paragraph("Third paragraph")
        doc.save(str(doc_path))
        return doc_path

    def _create_mapping(self) -> DocxTypstMapping:
        doc_content = "Multi Paragraph Document"
        full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]

        node1 = _create_node_id("doc.children.0.body.0", doc_content)
        entry1 = MappingEntry(node_id=node1, source_start=0, source_end=1)

        node2 = _create_node_id("doc.children.1.body.0", doc_content)
        entry2 = MappingEntry(node_id=node2, source_start=1, source_end=2)

        node3 = _create_node_id("doc.children.2.body.0", doc_content)
        entry3 = MappingEntry(node_id=node3, source_start=2, source_end=3)

        entries = {
            node1.canonical: entry1,
            node2.canonical: entry2,
            node3.canonical: entry3,
        }
        path_index = {
            "paragraphs/0": node1.canonical,
            "paragraphs/1": node2.canonical,
            "paragraphs/2": node3.canonical,
        }

        return DocxTypstMapping(
            version="1.0.0",
            full_hash=full_hash,
            entries=entries,
            path_index=path_index,
        )

    def test_locate_nodes_returns_correct_indices(self, multi_para_doc: Path) -> None:
        from typst_ir.polish.node_locator import locate_nodes

        mapping = self._create_mapping()
        result = locate_nodes(mapping, multi_para_doc)

        assert len(result) == 3
        for entry in mapping.entries.values():
            assert result[entry.node_id.canonical] == entry.source_start

    def test_locate_nodes_warns_when_source_end_exceeds_length(
        self, multi_para_doc: Path, caplog
    ) -> None:
        from typst_ir.polish.node_locator import locate_nodes

        doc_content = "Multi Paragraph Document"
        full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]
        node = _create_node_id("doc.children.0", doc_content)
        entry = MappingEntry(node_id=node, source_start=0, source_end=10)

        mapping = DocxTypstMapping(
            version="1.0.0",
            full_hash=full_hash,
            entries={node.canonical: entry},
            path_index={"paragraphs/0": node.canonical},
        )

        with caplog.at_level("WARNING"):
            result = locate_nodes(mapping, multi_para_doc)

        assert result[node.canonical] == 0
        assert "exceeds document paragraph count" in caplog.text

    def test_locate_nodes_raises_for_missing_file(self, tmp_path: Path) -> None:
        from typst_ir.polish.node_locator import locate_nodes

        mapping = DocxTypstMapping(full_hash="test")
        missing = tmp_path / "missing.docx"

        with pytest.raises(FileNotFoundError):
            locate_nodes(mapping, missing)

    def test_locate_nodes_raises_for_invalid_docx(self, tmp_path: Path) -> None:
        from typst_ir.polish.node_locator import locate_nodes

        mapping = DocxTypstMapping(full_hash="test")
        invalid = tmp_path / "invalid.docx"
        invalid.write_text("not a docx")

        with pytest.raises(ValueError):
            locate_nodes(mapping, invalid)


class TestValidateMapping:
    @pytest.fixture
    def table_doc(self, tmp_path: Path) -> Path:
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "table.docx"
        doc = Document()
        doc.add_paragraph("Intro")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Cell A"
        table.cell(0, 1).text = "Cell B"
        doc.save(str(doc_path))
        return doc_path

    def _create_valid_mapping(self) -> DocxTypstMapping:
        doc_content = "Table Document"
        full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]

        node1 = _create_node_id("doc.children.0.body.0", doc_content)
        entry1 = MappingEntry(node_id=node1, source_start=0, source_end=1)

        node2 = _create_node_id("doc.children.1.body.0", doc_content)
        entry2 = MappingEntry(node_id=node2, source_start=1, source_end=2)

        node3 = _create_node_id("doc.children.2.body.0", doc_content)
        entry3 = MappingEntry(node_id=node3, source_start=2, source_end=3)

        entries = {
            node1.canonical: entry1,
            node2.canonical: entry2,
            node3.canonical: entry3,
        }
        path_index = {
            "paragraphs/0": node1.canonical,
            "paragraphs/1": node2.canonical,
            "paragraphs/2": node3.canonical,
        }

        return DocxTypstMapping(
            version="1.0.0",
            full_hash=full_hash,
            entries=entries,
            path_index=path_index,
        )

    def test_validate_mapping_passes_for_valid_indices(self, table_doc: Path) -> None:
        from typst_ir.polish.node_locator import validate_mapping

        mapping = self._create_valid_mapping()
        validate_mapping(mapping, table_doc)

    def test_validate_mapping_raises_when_source_start_exceeds_bounds(
        self, table_doc: Path
    ) -> None:
        from typst_ir.polish.node_locator import validate_mapping

        doc_content = "Table Document"
        full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]
        node = _create_node_id("doc.children.0.body.0", doc_content)
        entry = MappingEntry(node_id=node, source_start=10, source_end=11)

        mapping = DocxTypstMapping(
            version="1.0.0",
            full_hash=full_hash,
            entries={node.canonical: entry},
            path_index={"paragraphs/0": node.canonical},
        )

        with pytest.raises(ValueError, match="source_start"):
            validate_mapping(mapping, table_doc)

    def test_validate_mapping_raises_when_source_end_exceeds_bounds(
        self, table_doc: Path
    ) -> None:
        from typst_ir.polish.node_locator import validate_mapping

        doc_content = "Table Document"
        full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]
        node = _create_node_id("doc.children.0.body.0", doc_content)
        entry = MappingEntry(node_id=node, source_start=0, source_end=10)

        mapping = DocxTypstMapping(
            version="1.0.0",
            full_hash=full_hash,
            entries={node.canonical: entry},
            path_index={"paragraphs/0": node.canonical},
        )

        with pytest.raises(ValueError, match="source_end"):
            validate_mapping(mapping, table_doc)

    def test_validate_mapping_raises_for_negative_index(self, table_doc: Path) -> None:
        from typst_ir.polish.node_locator import validate_mapping

        doc_content = "Table Document"
        full_hash = hashlib.sha256(doc_content.encode()).hexdigest()[:16]
        node = _create_node_id("doc.children.0.body.0", doc_content)
        entry = MappingEntry(node_id=node, source_start=-1, source_end=1)

        mapping = DocxTypstMapping(
            version="1.0.0",
            full_hash=full_hash,
            entries={node.canonical: entry},
            path_index={"paragraphs/0": node.canonical},
        )

        with pytest.raises(ValueError, match="negative index"):
            validate_mapping(mapping, table_doc)

    def test_validate_mapping_raises_for_missing_file(self, tmp_path: Path) -> None:
        from typst_ir.polish.node_locator import validate_mapping

        mapping = DocxTypstMapping(full_hash="test")
        missing = tmp_path / "missing.docx"

        with pytest.raises(FileNotFoundError):
            validate_mapping(mapping, missing)

    def test_validate_mapping_raises_for_invalid_docx(self, tmp_path: Path) -> None:
        from typst_ir.polish.node_locator import validate_mapping

        mapping = DocxTypstMapping(full_hash="test")
        invalid = tmp_path / "invalid.docx"
        invalid.write_text("not a docx")

        with pytest.raises(ValueError):
            validate_mapping(mapping, invalid)
