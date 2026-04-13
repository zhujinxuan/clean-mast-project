"""Tests for CLI commands in typst_ir.polish.cli.

Tests the generate-polish-ir and apply-polish-ir commands via the main
Click group registered in typst_ir.mapping.cli.
"""

from __future__ import annotations

import json
from pathlib import Path

from click.testing import CliRunner
from docx import Document
from docx.shared import Pt
from typst_ir import TypstDocument, TypstHeading, TypstParagraph, TypstSection
from typst_ir.mapping.cli import main


def _make_styled_docx(path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    doc.add_paragraph("Body text", style="Normal")
    doc.save(str(path))


def _make_config_toml(path: Path) -> None:
    path.write_text(
        '\n'
        '[styles."Normal"]\n'
        'font_name = "Arial"\n'
        'font_size = 14.0\n'
        '\n'
        '[styles."MissingStyle"]\n'
        'font_name = "Courier New"\n'
        '\n'
        '[renumber]\n'
        'headings = true\n'
        'figures = false\n'
        'tables = true\n',
        encoding="utf-8",
    )


def _make_ir_json(path: Path) -> None:
    ir = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Title"),
                body=[TypstParagraph(text="Body text")],
                subsections=[],
            ),
        ],
    )
    path.write_text(ir.model_dump_json(), encoding="utf-8")


class TestGeneratePolishIr:
    def test_generate_polish_ir_help(self) -> None:
        runner = CliRunner()
        result = runner.invoke(main, ["generate-polish-ir", "--help"])

        assert result.exit_code == 0
        assert "--docx" in result.output
        assert "--config" in result.output
        assert "--output-ir" in result.output

    def test_generate_polish_ir_missing_docx(self, tmp_path: Path) -> None:
        runner = CliRunner()
        config = tmp_path / "config.toml"
        _make_config_toml(config)
        result = runner.invoke(
            main,
            [
                "generate-polish-ir",
                "--config", str(config),
                "--output-ir", str(tmp_path / "out.json"),
            ],
        )
        assert result.exit_code != 0
        assert "--docx" in result.output or "Missing option" in result.output

    def test_generate_polish_ir_missing_config(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        result = runner.invoke(
            main,
            [
                "generate-polish-ir",
                "--docx", str(docx),
                "--output-ir", str(tmp_path / "out.json"),
            ],
        )
        assert result.exit_code != 0
        assert "--config" in result.output or "Missing option" in result.output

    def test_generate_polish_ir_missing_output_ir(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        config = tmp_path / "config.toml"
        _make_config_toml(config)
        result = runner.invoke(
            main,
            [
                "generate-polish-ir",
                "--docx", str(docx),
                "--config", str(config),
            ],
        )
        assert result.exit_code != 0
        assert "--output-ir" in result.output or "Missing option" in result.output

    def test_generate_polish_ir_success(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        config = tmp_path / "config.toml"
        _make_config_toml(config)
        output_ir = tmp_path / "out.json"

        result = runner.invoke(
            main,
            [
                "generate-polish-ir",
                "--docx", str(docx),
                "--config", str(config),
                "--output-ir", str(output_ir),
            ],
        )

        assert result.exit_code == 0, result.output
        assert output_ir.exists()
        data = json.loads(output_ir.read_text(encoding="utf-8"))
        assert data["version"] == "1.0.0"

        styles = {op["style_name"]: op for op in data["styles"]}
        assert "Normal" in styles
        assert styles["Normal"]["font_name"] == "Arial"
        assert styles["Normal"]["font_size"] == 14.0
        assert "MissingStyle" not in styles

        renumber_types = {op["element_type"] for op in data["renumber"]}
        assert renumber_types == {"heading", "table"}

    def test_generate_polish_ir_includes_line_spacing(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        config = tmp_path / "config.toml"
        config.write_text(
            '\n[styles."Normal"]\n'
            'font_name = "Arial"\n'
            'font_size = 14.0\n'
            'line_spacing = 1.5\n'
            '\n[renumber]\n'
            'headings = false\n'
            'figures = false\n'
            'tables = false\n',
            encoding="utf-8",
        )
        output_ir = tmp_path / "out.json"

        result = runner.invoke(
            main,
            [
                "generate-polish-ir",
                "--docx", str(docx),
                "--config", str(config),
                "--output-ir", str(output_ir),
            ],
        )

        assert result.exit_code == 0, result.output
        data = json.loads(output_ir.read_text(encoding="utf-8"))
        styles = {op["style_name"]: op for op in data["styles"]}
        assert "Normal" in styles
        assert styles["Normal"]["line_spacing"] == 1.5

    def test_generate_polish_ir_skips_empty_style_config(
        self, tmp_path: Path
    ) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        config = tmp_path / "config.toml"
        config.write_text(
            '\n[styles."Normal"]\n\n'
            '[renumber]\n'
            'headings = false\n'
            'figures = false\n'
            'tables = false\n',
            encoding="utf-8",
        )
        output_ir = tmp_path / "out.json"

        result = runner.invoke(
            main,
            [
                "generate-polish-ir",
                "--docx", str(docx),
                "--config", str(config),
                "--output-ir", str(output_ir),
            ],
        )

        assert result.exit_code == 0
        data = json.loads(output_ir.read_text(encoding="utf-8"))
        assert data["styles"] == []
        assert data["renumber"] == []


class TestApplyPolishIr:
    def test_apply_polish_ir_help(self) -> None:
        runner = CliRunner()
        result = runner.invoke(main, ["apply-polish-ir", "--help"])

        assert result.exit_code == 0
        assert "--docx" in result.output
        assert "--polish-ir" in result.output
        assert "--output-docx" in result.output
        assert "--ir-json" in result.output

    def test_apply_polish_ir_missing_docx(self, tmp_path: Path) -> None:
        runner = CliRunner()
        polish_ir = tmp_path / "polish.json"
        polish_ir.write_text(
            '{"version":"1.0.0","styles":[],"renumber":[]}',
            encoding="utf-8",
        )
        result = runner.invoke(
            main,
            [
                "apply-polish-ir",
                "--polish-ir", str(polish_ir),
                "--output-docx", str(tmp_path / "out.docx"),
            ],
        )
        assert result.exit_code != 0
        assert "--docx" in result.output or "Missing option" in result.output

    def test_apply_polish_ir_missing_polish_ir(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        result = runner.invoke(
            main,
            [
                "apply-polish-ir",
                "--docx", str(docx),
                "--output-docx", str(tmp_path / "out.docx"),
            ],
        )
        assert result.exit_code != 0
        assert "--polish-ir" in result.output or "Missing option" in result.output

    def test_apply_polish_ir_missing_output_docx(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        polish_ir = tmp_path / "polish.json"
        polish_ir.write_text(
            '{"version":"1.0.0","styles":[],"renumber":[]}',
            encoding="utf-8",
        )
        result = runner.invoke(
            main,
            [
                "apply-polish-ir",
                "--docx", str(docx),
                "--polish-ir", str(polish_ir),
            ],
        )
        assert result.exit_code != 0
        assert "--output-docx" in result.output or "Missing option" in result.output

    def test_apply_styles_success(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        polish_ir = tmp_path / "polish.json"
        polish_ir.write_text(
            json.dumps(
                {
                    "version": "1.0.0",
                    "styles": [
                        {
                            "style_name": "Normal",
                            "font_name": "Arial",
                            "font_size": 14.0,
                        }
                    ],
                    "renumber": [],
                }
            ),
            encoding="utf-8",
        )
        output_docx = tmp_path / "out.docx"

        result = runner.invoke(
            main,
            [
                "apply-polish-ir",
                "--docx", str(docx),
                "--polish-ir", str(polish_ir),
                "--output-docx", str(output_docx),
            ],
        )

        assert result.exit_code == 0, result.output
        assert output_docx.exists()

        out_doc = Document(str(output_docx))
        normal = out_doc.styles["Normal"]
        assert normal.font.name == "Arial"
        assert normal.font.size is not None
        assert normal.font.size.pt == 14.0

    def test_renumbering_requires_ir_json(self, tmp_path: Path) -> None:
        runner = CliRunner()
        docx = tmp_path / "test.docx"
        _make_styled_docx(docx)
        polish_ir = tmp_path / "polish.json"
        polish_ir.write_text(
            json.dumps(
                {
                    "version": "1.0.0",
                    "styles": [],
                    "renumber": [{"element_type": "heading"}],
                }
            ),
            encoding="utf-8",
        )
        output_docx = tmp_path / "out.docx"

        result = runner.invoke(
            main,
            [
                "apply-polish-ir",
                "--docx", str(docx),
                "--polish-ir", str(polish_ir),
                "--output-docx", str(output_docx),
            ],
        )

        assert result.exit_code != 0
        assert "--ir-json" in result.output

    def test_apply_renumbering_success(self, tmp_path: Path) -> None:
        runner = CliRunner()

        docx = tmp_path / "test.docx"
        doc = Document()
        doc.add_heading("Title", level=1)
        doc.save(str(docx))

        ir_json = tmp_path / "ir.json"
        _make_ir_json(ir_json)

        polish_ir = tmp_path / "polish.json"
        polish_ir.write_text(
            json.dumps(
                {
                    "version": "1.0.0",
                    "styles": [],
                    "renumber": [{"element_type": "heading"}],
                }
            ),
            encoding="utf-8",
        )
        output_docx = tmp_path / "out.docx"

        result = runner.invoke(
            main,
            [
                "apply-polish-ir",
                "--docx", str(docx),
                "--polish-ir", str(polish_ir),
                "--ir-json", str(ir_json),
                "--output-docx", str(output_docx),
            ],
        )

        assert result.exit_code == 0, result.output
        assert output_docx.exists()

        out_doc = Document(str(output_docx))
        assert out_doc.paragraphs[0].text == "1 Title"
