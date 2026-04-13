"""Tests for inner content extraction from DOCX documents."""

from pathlib import Path

import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from typst_ir.mapping.extractor import extract_inner_content


class TestExtractInnerContent:
    """Tests for extract_inner_content()."""

    @pytest.fixture
    def interleaved_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with paragraph, table, paragraph interleaving."""
        doc_path = tmp_path / "interleaved.docx"
        doc = Document()
        doc.add_paragraph("First paragraph")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "C"
        table.cell(1, 1).text = "D"
        doc.add_paragraph("Last paragraph")
        doc.save(str(doc_path))
        return doc_path

    def test_preserves_interleaving(self, interleaved_doc: Path) -> None:
        """extract_inner_content() preserves paragraph/table interleaving."""
        doc = Document(str(interleaved_doc))
        result = extract_inner_content(doc)

        assert len(result) == 3

        indexes = [idx for idx, _ in result]
        assert indexes == [0, 1, 2]

        types = [type(element) for _, element in result]
        assert types == [Paragraph, Table, Paragraph]

        assert isinstance(result[0][1], Paragraph)
        assert isinstance(result[1][1], Table)
        assert isinstance(result[2][1], Paragraph)
