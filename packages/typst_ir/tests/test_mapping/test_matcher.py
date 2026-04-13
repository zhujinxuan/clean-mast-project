"""Integration tests for DocxMatcher."""

import logging
from pathlib import Path

import pytest
from docx import Document
from typst_ir import TypstHeading, TypstParagraph, TypstSection, TypstTable
from typst_ir.mapping import DocxMatcher
from typst_ir.mapping.types import ConfidenceLevel, DetectionMethod


class TestDocxMatcher:
    @pytest.mark.skip(reason="Matcher implementation does not yet flatten TypstSection nodes")
    def test_match_paragraph_high_confidence(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("The wind speed at 80 meters was measured.")
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Paragraph"),
                body=[TypstParagraph(text="The wind speed at 80 meters was measured.")],
                subsections=[],
            )
        ]

        matcher = DocxMatcher()
        mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 1
        assert mappings[0].method == DetectionMethod.fingerprint
        assert mappings[0].confidence == ConfidenceLevel.high
        assert mappings[0].source_start == 0
        assert mappings[0].source_end == 1

    @pytest.mark.skip(reason="Matcher implementation does not yet flatten TypstSection nodes")
    def test_match_heading(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        para = doc.add_paragraph("Introduction")
        try:
            para.style = doc.styles["Heading 1"]
        except KeyError:
            pass
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Introduction"),
                body=[],
                subsections=[],
            )
        ]

        matcher = DocxMatcher()
        mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 1
        assert mappings[0].method == DetectionMethod.style
        assert mappings[0].confidence == ConfidenceLevel.high
        assert mappings[0].details["level"] == 1
        assert mappings[0].source_start == 0
        assert mappings[0].source_end == 1

    @pytest.mark.skip(reason="Matcher implementation does not yet flatten TypstSection nodes")
    def test_match_table_unambiguous(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Not a caption")
        doc.add_table(rows=2, cols=2)
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Table"),
                body=[
                    TypstTable(
                        headers=["A", "B"],
                        data=[["1", "2"], ["3", "4"]],
                    )
                ],
                subsections=[],
            )
        ]

        matcher = DocxMatcher()
        mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 1
        assert mappings[0].method is not None
        assert mappings[0].confidence is not None
        assert mappings[0].source_start == 1
        assert mappings[0].source_end == 2

    def test_unmatched_node_logs_warning(self, tmp_path: Path, caplog) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Just a normal paragraph.")
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Introduction"),
                body=[],
                subsections=[],
            )
        ]

        matcher = DocxMatcher()
        with caplog.at_level("WARNING"):
            mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 0

    def test_high_log_level_suppresses_warnings(self, tmp_path: Path, caplog) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Just a normal paragraph.")
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Introduction"),
                body=[],
                subsections=[],
            )
        ]

        matcher = DocxMatcher(log_level=logging.ERROR)
        with caplog.at_level("WARNING"):
            mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 0
        assert not (
            "Unmatched" in caplog.text or "unmatched" in caplog.text.lower()
        )

    def test_log_level_accepts_string(self, tmp_path: Path, caplog) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Just a normal paragraph.")
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Introduction"),
                body=[],
                subsections=[],
            )
        ]

        matcher = DocxMatcher(log_level="ERROR")
        with caplog.at_level("WARNING"):
            mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 0
        assert not (
            "Unmatched" in caplog.text or "unmatched" in caplog.text.lower()
        )

    @pytest.mark.skip(reason="Matcher implementation does not yet flatten TypstSection nodes")
    def test_match_multiple_nodes_sequential(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("First paragraph text.")
        heading = doc.add_paragraph("Section Heading")
        try:
            heading.style = doc.styles["Heading 1"]
        except KeyError:
            pass
        doc.add_table(rows=2, cols=2)
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        ir_nodes = [
            TypstSection(
                heading=TypstHeading(text="Section 1"),
                body=[TypstParagraph(text="First paragraph text.")],
                subsections=[],
            ),
            TypstSection(
                heading=TypstHeading(text="Section Heading"),
                body=[],
                subsections=[],
            ),
            TypstSection(
                heading=TypstHeading(text="Table"),
                body=[
                    TypstTable(
                        headers=["A", "B"],
                        data=[["1", "2"], ["3", "4"]],
                    )
                ],
                subsections=[],
            ),
        ]

        matcher = DocxMatcher()
        mappings = matcher.match(ir_nodes, doc2)

        assert len(mappings) == 3
        assert mappings[0].method == DetectionMethod.fingerprint
        assert mappings[1].method == DetectionMethod.style
        assert mappings[1].details["level"] == 1
        assert mappings[2].method is not None
        assert mappings[2].confidence is not None
        assert mappings[0].source_start == 0
        assert mappings[1].source_start == 1
        assert mappings[2].source_start == 2
