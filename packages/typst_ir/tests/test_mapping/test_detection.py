from pathlib import Path

import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from typst_ir import TypstFigure, TypstTable
from typst_ir.mapping.detection import (
    DetectionResult,
    detect_equation,
    detect_figures,
    detect_heading_level,
    detect_tables,
    match_figure,
    match_paragraph_fingerprint,
    match_paragraph_with_fallback,
    match_table,
)


class TestMatchParagraphFingerprint:
    def test_pure_text_exact_match(self) -> None:
        text = "The wind speed at 80 meters was measured using a calibrated anemometer."
        is_match, cov_a, cov_b = match_paragraph_fingerprint(text, text)
        assert is_match is True
        assert cov_a == 1.0
        assert cov_b == 1.0

    def test_reference_substitution_matches(self) -> None:
        typst = (
            "See Figure \\ref{fig1} for details. "
            "The wind speed data shows a clear diurnal pattern with maximum values "
            "occurring during afternoon hours."
        )
        docx = (
            "See Figure 1 for details. "
            "The wind speed data shows a clear diurnal pattern with maximum values "
            "occurring during afternoon hours."
        )
        is_match, cov_a, cov_b = match_paragraph_fingerprint(typst, docx)
        assert is_match is True
        assert cov_a >= 0.8
        assert cov_b >= 0.8

    def test_formatting_difference_normalizes(self) -> None:
        typst = "The  wind  speed  was  8.5  m/s  at  80m  height."
        docx = "The wind speed was 8.5 m/s at 80m height."
        is_match, cov_a, cov_b = match_paragraph_fingerprint(typst, docx)
        assert is_match is True
        assert cov_a == 1.0
        assert cov_b == 1.0

    def test_short_text_fails_threshold(self) -> None:
        typst = "See \\ref{fig1}."
        docx = "See Figure 1."
        is_match, cov_a, cov_b = match_paragraph_fingerprint(typst, docx)
        assert is_match is False
        assert cov_a == 0.0
        assert cov_b == 0.0

    def test_partial_match_below_threshold(self) -> None:
        typst = "Completely different text that has nothing in common."
        docx = "Another unrelated paragraph with different words entirely."
        is_match, cov_a, cov_b = match_paragraph_fingerprint(typst, docx)
        assert is_match is False
        assert cov_a < 0.8
        assert cov_b < 0.8


class TestMatchParagraphWithFallback:
    @pytest.fixture
    def make_para(self, tmp_path: Path):
        def _make(text: str, style_name: str = "Normal"):
            doc_path = tmp_path / "test.docx"
            doc = Document()
            para = doc.add_paragraph(text)
            try:
                para.style = doc.styles[style_name]
            except KeyError:
                pass
            doc.save(str(doc_path))
            doc2 = Document(str(doc_path))
            return doc2.paragraphs[0]

        return _make

    def test_fingerprint_high_confidence(self, make_para) -> None:
        typst = "The wind speed at 80 meters was measured using a calibrated anemometer."
        docx_para = make_para(typst, "Normal")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=0)
        assert result["method"] == "fingerprint"
        assert result["confidence"] == "high"
        assert result["coverage"] is not None
        assert result["coverage"][0] >= 0.8
        assert result["coverage"][1] >= 0.8

    def test_reference_substitution_high_confidence(self, make_para) -> None:
        typst = (
            "See Figure \\ref{fig1} for details. "
            "The wind speed data shows a clear diurnal pattern with maximum values."
        )
        docx = "See Figure 1 for details. The wind speed data shows a clear diurnal pattern with maximum values."
        docx_para = make_para(docx, "Normal")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=0)
        assert result["method"] == "fingerprint"
        assert result["confidence"] == "high"

    def test_style_fallback_short_heading(self, make_para) -> None:
        typst = "Short head"
        docx_para = make_para("Short head", "Heading 1")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=1)
        assert result["method"] == "style_heading"
        assert result["confidence"] == "medium"
        assert result["coverage"] is None

    def test_style_fallback_short_caption(self, make_para) -> None:
        typst = "A caption"
        docx_para = make_para("A caption", "Caption")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=2)
        assert result["method"] == "style_caption"
        assert result["confidence"] == "medium"
        assert result["coverage"] is None

    def test_hash_fallback_very_short_body(self, make_para) -> None:
        typst = "Hello world"
        docx_para = make_para("Hello world", "Body Text")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=3)
        assert result["method"] == "hash"
        assert result["confidence"] == "low"
        assert result["coverage"] is None

    def test_style_fallback_short_body_no_hash(self, make_para) -> None:
        typst = "Typst text"
        docx_para = make_para("Docx text", "Body Text")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=4)
        assert result["method"] == "style_body"
        assert result["confidence"] == "low"
        assert result["coverage"] is None

    def test_positional_fallback_long_mismatch(self, make_para) -> None:
        typst = "This is a long paragraph that will not match anything in the docx side."
        docx_para = make_para("A completely different long paragraph for docx content.", "Normal")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=0, docx_index=5)
        assert result["method"] == "positional"
        assert result["confidence"] == "low"
        assert result["coverage"] is None

    def test_passes_through_indices(self, make_para) -> None:
        typst = "Any text"
        docx_para = make_para("Any text", "Normal")
        result = match_paragraph_with_fallback(typst, docx_para, ir_index=7, docx_index=9)
        assert result["ir_index"] == 7
        assert result["docx_index"] == 9


class TestDetectionResult:
    def test_creation_and_access(self) -> None:
        result = DetectionResult(
            ir_node_id="node-123",
            ir_node_type="TypstParagraph",
            docx_index=5,
            method="fingerprint",
            confidence="high",
            details={"coverage": (0.9, 0.9)},
        )
        assert result.ir_node_id == "node-123"
        assert result.ir_node_type == "TypstParagraph"
        assert result.docx_index == 5
        assert result.method == "fingerprint"
        assert result.confidence == "high"
        assert result.details == {"coverage": (0.9, 0.9)}


class TestDetectHeadingLevel:
    @pytest.fixture
    def make_para(self, tmp_path: Path):
        def _make(text: str, style_name: str = "Normal"):
            doc_path = tmp_path / "test.docx"
            doc = Document()
            para = doc.add_paragraph(text)
            try:
                para.style = doc.styles[style_name]
            except KeyError:
                pass
            doc.save(str(doc_path))
            doc2 = Document(str(doc_path))
            return doc2.paragraphs[0]

        return _make

    def test_heading_1_style_detection(self, make_para) -> None:
        para = make_para("Introduction", "Heading 1")
        assert detect_heading_level(para) == 1

    def test_heading_2_style_detection(self, make_para) -> None:
        para = make_para("Methods", "Heading 2")
        assert detect_heading_level(para) == 2

    def test_heading_3_case_insensitive(self, make_para) -> None:
        para = make_para("Details", "heading 3")
        assert detect_heading_level(para) == 3

    def test_xml_fallback_outline_level(self, make_para) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn as _qn

        para = make_para("Custom heading", "Normal")
        pPr = para._element.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(_qn("w:val"), "2")
        pPr.append(outline)

        assert detect_heading_level(para) == 2

    def test_non_heading_returns_none(self, make_para) -> None:
        para = make_para("Just a normal paragraph.", "Normal")
        assert detect_heading_level(para) is None


class TestDetectTables:
    def test_detects_tables_in_order_with_interleaving(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_table(rows=1, cols=2)
        doc.add_paragraph("Between tables")
        doc.add_table(rows=2, cols=3)
        doc.add_paragraph("Last paragraph")
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        tables = detect_tables(doc2)

        assert len(tables) == 2
        assert tables[0][0] == 1
        assert tables[1][0] == 3


class TestMatchTable:
    def test_high_confidence_structure_and_caption(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        caption_para = doc.add_paragraph("Wind resource parameters")
        try:
            caption_para.style = doc.styles["Caption"]
        except KeyError:
            pass
        doc.add_table(rows=3, cols=2)
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        contents = list(doc2.iter_inner_content())
        docx_table = contents[1]
        assert isinstance(docx_table, Table)

        ir_table = TypstTable(
            headers=["Parameter", "Value"],
            data=[["Mean wind speed", "7.5"], ["Weibull k", "2.1"]],
            caption="Wind resource parameters",
        )

        result = match_table(ir_table, docx_table, contents, table_index=1)

        assert result["structure"] is True
        assert result["caption"] is True
        assert result["confidence"] == "high"
        assert result["ir_dims"] == (3, 2)
        assert result["docx_dims"] == (3, 2)

    def test_medium_confidence_structure_only(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Not a caption")
        doc.add_table(rows=2, cols=2)
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        contents = list(doc2.iter_inner_content())
        docx_table = contents[1]
        assert isinstance(docx_table, Table)

        ir_table = TypstTable(
            headers=["A", "B"],
            data=[["1", "2"]],
            caption="Some caption",
        )

        result = match_table(ir_table, docx_table, contents, table_index=1)

        assert result["structure"] is True
        assert result["caption"] is None
        assert result["confidence"] == "medium"

    def test_low_confidence_no_match(self, tmp_path: Path) -> None:
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Not a caption")
        doc.add_table(rows=1, cols=1)
        doc.save(str(doc_path))

        doc2 = Document(str(doc_path))
        contents = list(doc2.iter_inner_content())
        docx_table = contents[1]
        assert isinstance(docx_table, Table)

        ir_table = TypstTable(
            headers=["A", "B", "C"],
            data=[["1", "2", "3"], ["4", "5", "6"]],
            caption="Different caption",
        )

        result = match_table(ir_table, docx_table, contents, table_index=1)

        assert result["structure"] is False
        assert result["caption"] is None
        assert result["confidence"] == "low"


class TestDetectFigures:
    @pytest.fixture
    def make_docx_figure(self, tmp_path: Path):
        from docx.oxml import OxmlElement

        def _make(*paragraphs: tuple[str, str | None, bool]):
            doc_path = tmp_path / "test.docx"
            doc = Document()
            for text, style_name, has_drawing in paragraphs:
                para = doc.add_paragraph(text)
                if style_name:
                    try:
                        para.style = doc.styles[style_name]
                    except KeyError:
                        pass
                if has_drawing:
                    run = para.add_run()
                    drawing = OxmlElement("w:drawing")
                    run._r.append(drawing)
            doc.save(str(doc_path))
            return Document(str(doc_path))

        return _make

    def test_detect_figures_finds_drawing_paragraphs(self, make_docx_figure) -> None:
        doc = make_docx_figure(
            ("First paragraph", "Normal", False),
            ("Figure content", "Normal", True),
            ("Last paragraph", "Normal", False),
        )
        figures = detect_figures(doc)

        assert len(figures) == 1
        assert figures[0]["index"] == 1
        assert figures[0]["has_drawing"] is True

    def test_detect_figures_associates_captions(self, make_docx_figure) -> None:
        doc = make_docx_figure(
            ("Figure content", "Normal", True),
            ("Wind speed diagram", "Caption", False),
        )
        figures = detect_figures(doc)

        assert len(figures) == 1
        assert figures[0]["index"] == 0
        assert figures[0]["caption"] == "Wind speed diagram"
        assert figures[0]["caption_index"] == 1


class TestMatchFigure:
    @pytest.fixture
    def make_docx_figure(self, tmp_path: Path):
        from docx.oxml import OxmlElement

        def _make(*paragraphs: tuple[str, str | None, bool]):
            doc_path = tmp_path / "test.docx"
            doc = Document()
            for text, style_name, has_drawing in paragraphs:
                para = doc.add_paragraph(text)
                if style_name:
                    try:
                        para.style = doc.styles[style_name]
                    except KeyError:
                        pass
                if has_drawing:
                    run = para.add_run()
                    drawing = OxmlElement("w:drawing")
                    run._r.append(drawing)
            doc.save(str(doc_path))
            doc2 = Document(str(doc_path))
            return list(doc2.iter_inner_content()), doc2

        return _make

    def test_high_confidence_drawing_and_caption(self, make_docx_figure) -> None:
        contents, doc = make_docx_figure(
            ("Figure content", "Normal", True),
            ("Wind speed diagram at 80 meters height", "Caption", False),
        )
        docx_para = contents[0]
        assert isinstance(docx_para, Paragraph)

        ir_figure = TypstFigure(
            content="image.png",
            caption="Wind speed diagram at 80 meters height",
            kind="image",
        )
        result = match_figure(ir_figure, docx_para, contents, para_index=0)

        assert result["drawing"] is True
        assert result["caption"] is True
        assert result["confidence"] == "high"

    def test_medium_confidence_drawing_only(self, make_docx_figure) -> None:
        contents, doc = make_docx_figure(
            ("Figure content", "Normal", True),
            ("Some regular paragraph", "Normal", False),
        )
        docx_para = contents[0]
        assert isinstance(docx_para, Paragraph)

        ir_figure = TypstFigure(
            content="image.png",
            caption="Wind speed diagram",
            kind="image",
        )
        result = match_figure(ir_figure, docx_para, contents, para_index=0)

        assert result["drawing"] is True
        assert result["caption"] is None
        assert result["confidence"] == "medium"

    def test_low_confidence_no_match(self, make_docx_figure) -> None:
        contents, doc = make_docx_figure(
            ("Just a normal paragraph", "Normal", False),
            ("Another regular paragraph", "Normal", False),
        )
        docx_para = contents[0]
        assert isinstance(docx_para, Paragraph)

        ir_figure = TypstFigure(
            content="image.png",
            caption="Wind speed diagram",
            kind="image",
        )
        result = match_figure(ir_figure, docx_para, contents, para_index=0)

        assert result["drawing"] is False
        assert result["caption"] is None
        assert result["confidence"] == "low"


class TestDetectEquation:
    @pytest.fixture
    def make_equation_para(self, tmp_path: Path):
        from docx.oxml import OxmlElement

        def _make(tag_name: str | None):
            doc_path = tmp_path / "test.docx"
            doc = Document()
            para = doc.add_paragraph("Some text")
            if tag_name is not None:
                run = para.add_run()
                elem = OxmlElement(tag_name)
                run._r.append(elem)
            doc.save(str(doc_path))
            doc2 = Document(str(doc_path))
            return doc2.paragraphs[0]

        return _make

    def test_detects_display_equation(self, make_equation_para) -> None:
        para = make_equation_para("m:oMathPara")
        assert detect_equation(para) == "display"

    def test_detects_inline_equation(self, make_equation_para) -> None:
        para = make_equation_para("m:oMath")
        assert detect_equation(para) == "inline"

    def test_non_equation_returns_none(self, make_equation_para) -> None:
        para = make_equation_para(None)
        assert detect_equation(para) is None
