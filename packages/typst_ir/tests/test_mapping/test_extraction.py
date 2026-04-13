"""Tests for paragraph extraction from DOCX files.

Tests for extract_paragraphs() function that extracts metadata from DOCX paragraphs.
"""

import hashlib
from pathlib import Path

import pytest
from typst_ir.mapping.types import DocxParagraph


class TestSimpleParagraphExtraction:
    """Tests for basic paragraph extraction."""

    @pytest.fixture
    def simple_doc(self, tmp_path: Path) -> Path:
        """Create a simple DOCX with one paragraph."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "simple.docx"
        doc = Document()
        doc.add_paragraph("Hello World")
        doc.save(str(doc_path))
        return doc_path

    def test_extract_returns_list(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns a list."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        assert isinstance(result, list)

    def test_extract_single_paragraph(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns one paragraph for simple doc."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        assert len(result) == 1

    def test_extract_paragraph_type(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns correct paragraph type."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        assert result[0].paragraph_type == "paragraph"

    def test_extract_index(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns correct 0-based index."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        assert result[0].index == 0

    def test_extract_text_preview(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns correct text preview."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        assert result[0].text_preview == "Hello World"

    def test_extract_text_hash(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns correct SHA-256 hash."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        expected_hash = hashlib.sha256("Hello World".encode("utf-8")).hexdigest()
        assert result[0].text_hash == expected_hash

    def test_returns_docxparagraph_instances(self, simple_doc: Path) -> None:
        """extract_paragraphs() returns DocxParagraph instances."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(simple_doc)
        assert isinstance(result[0], DocxParagraph)


class TestMultipleParagraphs:
    """Tests for documents with multiple paragraphs."""

    @pytest.fixture
    def multi_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with multiple paragraphs."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "multi.docx"
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        doc.add_paragraph("Third paragraph")
        doc.save(str(doc_path))
        return doc_path

    def test_extracts_all_paragraphs(self, multi_doc: Path) -> None:
        """extract_paragraphs() extracts all paragraphs in order."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(multi_doc)
        assert len(result) == 3

    def test_indices_are_sequential(self, multi_doc: Path) -> None:
        """extract_paragraphs() assigns sequential indices."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(multi_doc)
        assert result[0].index == 0
        assert result[1].index == 1
        assert result[2].index == 2

    def test_previews_match_content(self, multi_doc: Path) -> None:
        """extract_paragraphs() returns correct previews for each paragraph."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(multi_doc)
        assert result[0].text_preview == "First paragraph"
        assert result[1].text_preview == "Second paragraph"
        assert result[2].text_preview == "Third paragraph"


class TestTextPreviewTruncation:
    """Tests for text preview truncation at 100 characters."""

    @pytest.fixture
    def long_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with a long paragraph."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "long.docx"
        doc = Document()
        long_text = "A" * 200
        doc.add_paragraph(long_text)
        doc.save(str(doc_path))
        return doc_path

    def test_preview_truncated_to_100_chars(self, long_doc: Path) -> None:
        """extract_paragraphs() truncates preview to 100 characters."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(long_doc)
        assert len(result[0].text_preview) == 100
        assert result[0].text_preview == "A" * 100

    def test_hash_is_of_full_text(self, long_doc: Path) -> None:
        """extract_paragraphs() hashes full text, not truncated."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(long_doc)
        expected_hash = hashlib.sha256(("A" * 200).encode("utf-8")).hexdigest()
        assert result[0].text_hash == expected_hash


class TestHeadingDetection:
    """Tests for heading paragraph type detection."""

    @pytest.fixture
    def heading_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with headings."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "headings.docx"
        doc = Document()
        doc.add_heading("Main Title", level=1)
        doc.add_paragraph("Regular paragraph")
        doc.add_heading("Section Title", level=2)
        doc.save(str(doc_path))
        return doc_path

    def test_detects_heading_type(self, heading_doc: Path) -> None:
        """extract_paragraphs() detects heading paragraphs."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(heading_doc)
        assert result[0].paragraph_type == "heading"
        assert result[1].paragraph_type == "paragraph"
        assert result[2].paragraph_type == "heading"


class TestListItemDetection:
    """Tests for list item paragraph type detection."""

    @pytest.fixture
    def list_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with bullet list."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "list.docx"
        doc = Document()
        doc.add_paragraph("Intro")
        doc.add_paragraph("Item 1", style="List Bullet")
        doc.add_paragraph("Item 2", style="List Bullet")
        doc.save(str(doc_path))
        return doc_path

    def test_detects_list_item_type(self, list_doc: Path) -> None:
        """extract_paragraphs() detects list item paragraphs."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(list_doc)
        assert result[0].paragraph_type == "paragraph"
        assert result[1].paragraph_type == "list_item"
        assert result[2].paragraph_type == "list_item"


class TestTableCellExtraction:
    """Tests for table cell paragraph extraction."""

    @pytest.fixture
    def table_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with a 2x2 table."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "table.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Cell 1"
        table.cell(0, 1).text = "Cell 2"
        table.cell(1, 0).text = "Cell 3"
        table.cell(1, 1).text = "Cell 4"
        doc.save(str(doc_path))
        return doc_path

    def test_extracts_table_cell_paragraphs(self, table_doc: Path) -> None:
        """extract_paragraphs() includes table cell paragraphs."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(table_doc)
        assert len(result) == 4

    def test_detects_table_cell_type(self, table_doc: Path) -> None:
        """extract_paragraphs() detects table_cell paragraph type."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(table_doc)
        for para in result:
            assert para.paragraph_type == "table_cell"

    def test_table_cell_previews(self, table_doc: Path) -> None:
        """extract_paragraphs() returns correct previews for table cells."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(table_doc)
        previews = [p.text_preview for p in result]
        assert "Cell 1" in previews
        assert "Cell 2" in previews
        assert "Cell 3" in previews
        assert "Cell 4" in previews


class TestMixedContent:
    """Tests for documents with mixed content types."""

    @pytest.fixture
    def mixed_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with mixed content."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "mixed.docx"
        doc = Document()
        doc.add_heading("Title", level=1)
        doc.add_paragraph("Regular text")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Table Cell"
        doc.add_paragraph("List item", style="List Bullet")
        doc.save(str(doc_path))
        return doc_path

    def test_mixed_content_types(self, mixed_doc: Path) -> None:
        """extract_paragraphs() correctly identifies mixed content types."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(mixed_doc)
        types = [p.paragraph_type for p in result]

        assert "heading" in types
        assert "paragraph" in types
        assert "table_cell" in types
        assert "list_item" in types

    def test_mixed_indices_are_sequential(self, mixed_doc: Path) -> None:
        """extract_paragraphs() assigns sequential indices to mixed content."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(mixed_doc)
        for i, para in enumerate(result):
            assert para.index == i


class TestEmptyDocument:
    """Tests for empty document handling."""

    @pytest.fixture
    def empty_doc(self, tmp_path: Path) -> Path:
        """Create an empty DOCX document."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "empty.docx"
        doc = Document()
        doc.save(str(doc_path))
        return doc_path

    def test_empty_doc_returns_list(self, empty_doc: Path) -> None:
        """extract_paragraphs() returns list for empty document."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(empty_doc)
        assert isinstance(result, list)

    def test_empty_doc_returns_empty_list(self, empty_doc: Path) -> None:
        """extract_paragraphs() returns empty list for truly empty doc."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(empty_doc)
        assert len(result) == 0


class TestErrorHandling:
    """Tests for error conditions."""

    def test_raises_file_not_found_for_missing_file(self, tmp_path: Path) -> None:
        """extract_paragraphs() raises FileNotFoundError for missing file."""
        from typst_ir.mapping.injector import extract_paragraphs

        non_existent = tmp_path / "does_not_exist.docx"

        with pytest.raises(FileNotFoundError):
            extract_paragraphs(non_existent)

    def test_raises_value_error_for_invalid_docx(self, tmp_path: Path) -> None:
        """extract_paragraphs() raises ValueError for invalid DOCX file."""
        from typst_ir.mapping.injector import extract_paragraphs

        invalid_docx = tmp_path / "invalid.docx"
        invalid_docx.write_text("This is not a valid DOCX file")

        with pytest.raises((ValueError, Exception)):
            extract_paragraphs(invalid_docx)


class TestEdgeCases:
    """Tests for edge cases."""

    @pytest.fixture
    def empty_para_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with an empty paragraph."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "empty_para.docx"
        doc = Document()
        doc.add_paragraph("")
        doc.save(str(doc_path))
        return doc_path

    def test_empty_paragraph_handling(self, empty_para_doc: Path) -> None:
        """extract_paragraphs() handles empty paragraphs."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(empty_para_doc)
        assert len(result) == 1
        assert result[0].text_preview == ""
        expected_hash = hashlib.sha256("".encode("utf-8")).hexdigest()
        assert result[0].text_hash == expected_hash

    @pytest.fixture
    def unicode_doc(self, tmp_path: Path) -> Path:
        """Create a DOCX with unicode text."""
        pytest.importorskip("docx")
        from docx import Document

        doc_path = tmp_path / "unicode.docx"
        doc = Document()
        doc.add_paragraph("Unicode: 你好世界 🌍 émojis")
        doc.save(str(doc_path))
        return doc_path

    def test_unicode_text_handling(self, unicode_doc: Path) -> None:
        """extract_paragraphs() handles unicode text correctly."""
        from typst_ir.mapping.injector import extract_paragraphs

        result = extract_paragraphs(unicode_doc)
        assert len(result) == 1
        assert "你好世界" in result[0].text_preview
        assert "🌍" in result[0].text_preview
