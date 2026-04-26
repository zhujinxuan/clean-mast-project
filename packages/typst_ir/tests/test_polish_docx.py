"""Tests for polish_docx pipeline.

Covers the full pipeline: extract ranges, build block type map,
create ChangeStyles, and apply to DOCX.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from typst_ir.nodes import (
    TypstDocument,
    TypstHeading,
    TypstParagraph,
    TypstSection,
    TypstTable,
)
from typst_ir.polish_docx import polish_docx


def _ensure_style(doc: Document, style_name: str, style_type: WD_STYLE_TYPE) -> object:
    try:
        return doc.styles[style_name]
    except KeyError:
        return doc.styles.add_style(style_name, style_type)


def test_polish_docx_with_paragraphs(tmp_path: Path) -> None:
    """Test polishing a DOCX with paragraph markers."""
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]

    start = doc.add_paragraph("irstart:section0body0")
    for run in start.runs:
        run.style = verbatim

    content = doc.add_paragraph("This is content")

    end = doc.add_paragraph("irend:section0body0")
    for run in end.runs:
        run.style = verbatim

    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Section"),
                body=[TypstParagraph(text="This is content")],
                subsections=[],
            )
        ]
    )
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"

    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.paragraphs) == 1
    assert result_doc.paragraphs[0].text == "This is content"
    assert result_doc.paragraphs[0].style.name == "typst-ir-Paragraph"


def test_polish_docx_with_table(tmp_path: Path) -> None:
    """Test polishing a DOCX with table markers."""
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]

    start = doc.add_paragraph("irstart:section0body0table0")
    for run in start.runs:
        run.style = verbatim

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"

    end = doc.add_paragraph("irend:section0body0table0")
    for run in end.runs:
        run.style = verbatim

    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Section"),
                body=[
                    TypstTable(
                        headers=["Col1", "Col2"],
                        data=[["A", "B"], ["C", "D"]],
                        caption=None,
                    )
                ],
                subsections=[],
            )
        ]
    )
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"

    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.paragraphs) == 0
    assert len(result_doc.tables) == 1

    table = result_doc.tables[0]
    for row in table.rows:
        for cell in row.cells:
            assert cell.paragraphs[0].style.name == "typst-ir-TableContent"


def test_polish_docx_with_heading(tmp_path: Path) -> None:
    """Test polishing a DOCX with heading."""
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]

    start = doc.add_paragraph("irstart:section0heading0")
    for run in start.runs:
        run.style = verbatim

    heading = doc.add_paragraph("Section Heading")

    end = doc.add_paragraph("irend:section0heading0")
    for run in end.runs:
        run.style = verbatim

    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Section Heading"),
                body=[],
                subsections=[],
            )
        ]
    )
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"

    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.paragraphs) == 1
    assert result_doc.paragraphs[0].text == "Section Heading"
    assert result_doc.paragraphs[0].style.name == "typst-ir-Heading1"


def test_polish_docx_empty_document(tmp_path: Path) -> None:
    """Test polishing an empty DOCX."""
    doc = Document()
    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(children=[])
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"

    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.paragraphs) == 0


def test_polish_docx_table_format_single_table_has_borders_and_alignment(
    tmp_path: Path,
) -> None:
    """A document with 1 table produces table borders and center alignment."""
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]

    start = doc.add_paragraph("irstart:section0body0table0")
    for run in start.runs:
        run.style = verbatim

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"

    end = doc.add_paragraph("irend:section0body0table0")
    for run in end.runs:
        run.style = verbatim

    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Section"),
                body=[
                    TypstTable(
                        headers=["Col1", "Col2"],
                        data=[["A", "B"], ["C", "D"]],
                        caption=None,
                    )
                ],
                subsections=[],
            )
        ]
    )
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"
    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.tables) == 1

    table = result_doc.tables[0]
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    assert borders is not None
    assert len(borders) == 6

    jc = tbl_pr.find(qn("w:jc"))
    assert jc is not None
    assert jc.get(qn("w:val")) == "center"

    for row in table.rows:
        for cell in row.cells:
            assert cell.paragraphs[0].style.name == "typst-ir-TableContent"


def test_polish_docx_table_format_two_tables(tmp_path: Path) -> None:
    """A document with 2 tables produces borders on both tables."""
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]

    start1 = doc.add_paragraph("irstart:section0body0table0")
    for run in start1.runs:
        run.style = verbatim

    table1 = doc.add_table(rows=1, cols=2)
    table1.cell(0, 0).text = "A"
    table1.cell(0, 1).text = "B"

    end1 = doc.add_paragraph("irend:section0body0table0")
    for run in end1.runs:
        run.style = verbatim

    start2 = doc.add_paragraph("irstart:section0body1table1")
    for run in start2.runs:
        run.style = verbatim

    table2 = doc.add_table(rows=1, cols=2)
    table2.cell(0, 0).text = "X"
    table2.cell(0, 1).text = "Y"

    end2 = doc.add_paragraph("irend:section0body1table1")
    for run in end2.runs:
        run.style = verbatim

    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Section"),
                body=[
                    TypstTable(
                        headers=["Col1", "Col2"],
                        data=[["A", "B"]],
                        caption=None,
                    ),
                    TypstTable(
                        headers=["Col3", "Col4"],
                        data=[["X", "Y"]],
                        caption=None,
                    ),
                ],
                subsections=[],
            )
        ]
    )
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"
    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.tables) == 2

    for table in result_doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(qn("w:tblBorders"))
        assert borders is not None
        assert len(borders) == 6


def test_polish_docx_table_format_no_tables(tmp_path: Path) -> None:
    """A document with 0 tables produces no table borders."""
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]

    start = doc.add_paragraph("irstart:section0body0")
    for run in start.runs:
        run.style = verbatim

    content = doc.add_paragraph("Just a paragraph")

    end = doc.add_paragraph("irend:section0body0")
    for run in end.runs:
        run.style = verbatim

    docx_path = tmp_path / "input.docx"
    doc.save(str(docx_path))

    ir_doc = TypstDocument(
        children=[
            TypstSection(
                heading=TypstHeading(text="Section"),
                body=[TypstParagraph(text="Just a paragraph")],
                subsections=[],
            )
        ]
    )
    ir_path = tmp_path / "ir.json"
    ir_path.write_text(ir_doc.model_dump_json(), encoding="utf-8")

    output_path = tmp_path / "output.docx"
    polish_docx(docx_path, ir_path, output_path)

    result_doc = Document(str(output_path))
    assert len(result_doc.tables) == 0
    assert len(result_doc.paragraphs) == 1

    body = result_doc.element.body
    all_tbl_borders = body.findall(qn("w:tblBorders"))
    assert len(all_tbl_borders) == 0
