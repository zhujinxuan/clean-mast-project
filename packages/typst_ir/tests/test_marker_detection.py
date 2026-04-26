"""TDD tests for IR marker detection.

Tests parse_marker and is_marker_paragraph functions.
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from typst_ir.marker_detection import MARKER_PATTERN, is_marker_paragraph, parse_marker


class TestParseMarker:
    """Tests for parse_marker function."""

    def test_valid_irstart_marker(self) -> None:
        result = parse_marker("irstart:section0body0")
        assert result == ("irstart", "section0body0")

    def test_valid_irend_marker(self) -> None:
        result = parse_marker("irend:section0body2table0")
        assert result == ("irend", "section0body2table0")

    def test_empty_text(self) -> None:
        assert parse_marker("") is None

    def test_plain_text_no_marker(self) -> None:
        assert parse_marker("hello world") is None

    def test_partial_marker_no_colon(self) -> None:
        assert parse_marker("irstartsection0body0") is None

    def test_partial_marker_unknown_kind(self) -> None:
        assert parse_marker("irunknown:section0body0") is None

    def test_missing_block_id(self) -> None:
        assert parse_marker("irstart:") is None

    def test_whitespace_before_marker(self) -> None:
        assert parse_marker(" irstart:section0body0") is None

    def test_whitespace_after_kind(self) -> None:
        assert parse_marker("irstart :section0body0") is None


class TestIsMarkerParagraph:
    """Tests for is_marker_paragraph function."""

    @pytest.fixture
    def doc(self) -> Document:
        d = Document()
        d.styles.add_style("VerbatimChar", WD_STYLE_TYPE.CHARACTER)
        d.styles.add_style("OtherChar", WD_STYLE_TYPE.CHARACTER)
        return d

    def test_single_run_verbatimchar_valid_marker(self, doc: Document) -> None:
        p = doc.add_paragraph()
        run = p.add_run("irstart:section0body0")
        run.style = "VerbatimChar"
        assert is_marker_paragraph(p) is True

    def test_single_run_verbatimchar_irend_marker(self, doc: Document) -> None:
        p = doc.add_paragraph()
        run = p.add_run("irend:section0body0")
        run.style = "VerbatimChar"
        assert is_marker_paragraph(p) is True

    def test_multiple_runs(self, doc: Document) -> None:
        p = doc.add_paragraph()
        run1 = p.add_run("irstart:section0body0")
        run1.style = "VerbatimChar"
        run2 = p.add_run("extra")
        run2.style = "VerbatimChar"
        assert is_marker_paragraph(p) is False

    def test_no_runs(self, doc: Document) -> None:
        p = doc.add_paragraph()
        assert is_marker_paragraph(p) is False

    def test_wrong_style(self, doc: Document) -> None:
        p = doc.add_paragraph()
        run = p.add_run("irstart:section0body0")
        run.style = "OtherChar"
        assert is_marker_paragraph(p) is False

    def test_no_style(self, doc: Document) -> None:
        p = doc.add_paragraph()
        p.add_run("irstart:section0body0")
        assert is_marker_paragraph(p) is False

    def test_invalid_text_with_verbatimchar(self, doc: Document) -> None:
        p = doc.add_paragraph()
        run = p.add_run("not a marker")
        run.style = "VerbatimChar"
        assert is_marker_paragraph(p) is False

    def test_partial_marker_with_verbatimchar(self, doc: Document) -> None:
        p = doc.add_paragraph()
        run = p.add_run("irstart:")
        run.style = "VerbatimChar"
        assert is_marker_paragraph(p) is False


class TestMarkerPattern:
    """Direct tests for the MARKER_PATTERN regex."""

    def test_matches_irstart(self) -> None:
        match = MARKER_PATTERN.match("irstart:section0body0")
        assert match is not None
        assert match.group(1) == "irstart"
        assert match.group(2) == "section0body0"

    def test_matches_irend(self) -> None:
        match = MARKER_PATTERN.match("irend:section0body2table0")
        assert match is not None
        assert match.group(1) == "irend"
        assert match.group(2) == "section0body2table0"

    def test_no_match_unknown_kind(self) -> None:
        assert MARKER_PATTERN.match("irunknown:foo") is None

    def test_no_match_missing_id(self) -> None:
        assert MARKER_PATTERN.match("irstart:") is None
