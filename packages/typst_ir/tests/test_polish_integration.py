"""Integration test for polish_docx using real report_draft.docx."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from typst_ir.marker_detection import is_marker_paragraph
from typst_ir.polish_docx import polish_docx

# Resolve test data relative to workspace root (3 levels up from this file)
_WORKSPACE_ROOT = Path(__file__).resolve().parents[3]
_DOCX_PATH = _WORKSPACE_ROOT / "test-data/typst-ir/kp-draft-docx/work/rep-ts/report_draft.docx"
_IR_JSON_PATH = _WORKSPACE_ROOT / "test-data/typst-ir/kp-draft-docx/work/rep-ts/report.typ.json"


class TestPolishIntegrationRealData:
    def test_full_pipeline_removes_all_markers(self, tmp_path: Path) -> None:
        docx_path = _DOCX_PATH
        ir_json_path = _IR_JSON_PATH
        output_path = tmp_path / "polished_report.docx"

        assert docx_path.exists()
        assert ir_json_path.exists()

        input_doc = Document(str(docx_path))
        input_markers = [p for p in input_doc.paragraphs if is_marker_paragraph(p)]
        assert len(input_markers) > 0

        polish_docx(docx_path, ir_json_path, output_path)

        assert output_path.exists()

        output_doc = Document(str(output_path))
        remaining_markers = [p for p in output_doc.paragraphs if is_marker_paragraph(p)]
        assert len(remaining_markers) == 0

    def test_output_docx_is_valid_xml(self, tmp_path: Path) -> None:
        docx_path = _DOCX_PATH
        ir_json_path = _IR_JSON_PATH
        output_path = tmp_path / "polished_report.docx"

        polish_docx(docx_path, ir_json_path, output_path)

        assert zipfile.is_zipfile(output_path)
        with zipfile.ZipFile(output_path, "r") as zf:
            assert "word/document.xml" in zf.namelist()
            document_xml = zf.read("word/document.xml")
            assert len(document_xml) > 0
            content = document_xml.decode("utf-8", errors="replace")
            assert "<w:document" in content or "<?xml" in content

    def test_output_has_content_paragraphs(self, tmp_path: Path) -> None:
        docx_path = _DOCX_PATH
        ir_json_path = _IR_JSON_PATH
        output_path = tmp_path / "polished_report.docx"

        polish_docx(docx_path, ir_json_path, output_path)

        output_doc = Document(str(output_path))
        content_paragraphs = [p for p in output_doc.paragraphs if p.text.strip()]
        assert len(content_paragraphs) > 0

    def test_all_marker_pairs_removed(self, tmp_path: Path) -> None:
        docx_path = _DOCX_PATH
        ir_json_path = _IR_JSON_PATH
        output_path = tmp_path / "polished_report.docx"

        input_doc = Document(str(docx_path))
        input_markers = [p for p in input_doc.paragraphs if is_marker_paragraph(p)]

        polish_docx(docx_path, ir_json_path, output_path)

        output_doc = Document(str(output_path))
        output_markers = [p for p in output_doc.paragraphs if is_marker_paragraph(p)]

        assert len(output_markers) == 0, (
            f"Expected 0 markers in output, found {len(output_markers)}. "
            f"Input had {len(input_markers)} markers."
        )

    def test_input_has_expected_marker_count(self, tmp_path: Path) -> None:
        docx_path = _DOCX_PATH

        input_doc = Document(str(docx_path))
        input_markers = [p for p in input_doc.paragraphs if is_marker_paragraph(p)]

        assert len(input_markers) == 388, (
            f"Expected 388 marker paragraphs (194 pairs), found {len(input_markers)}"
        )
