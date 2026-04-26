from __future__ import annotations

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from typst_ir.marker_detection import is_marker_paragraph
from typst_ir.polish_executor import (
    _apply_style_properties,
    apply_change_style,
    apply_create_style,
    apply_edit_style,
    apply_table_format,
)
from typst_ir.polish_ir import (
    ApplyTableFormat,
    ChangeStyle,
    CreateStyle,
    EditStyle,
    TableFormat,
)
from typst_ir.polish_styles import LineSpacing, StyleProperties


def _ensure_style(doc: Document, style_name: str, style_type: WD_STYLE_TYPE) -> object:
    try:
        return doc.styles[style_name]
    except KeyError:
        return doc.styles.add_style(style_name, style_type)


def test_apply_change_style_to_paragraphs() -> None:
    doc = Document()
    target_style = doc.styles.add_style("TestStyle", WD_STYLE_TYPE.PARAGRAPH)

    p0 = doc.add_paragraph("Paragraph 0")
    p1 = doc.add_paragraph("Paragraph 1")
    p2 = doc.add_paragraph("Paragraph 2")

    ir = ChangeStyle(
        path="section0body0",
        target_style="TestStyle",
        paragraph_indices=(0, 2),
    )

    apply_change_style(doc, ir)

    assert p0.style == target_style
    assert p1.style != target_style
    assert p2.style == target_style


def test_apply_change_style_to_table_cells() -> None:
    doc = Document()
    target_style = doc.styles.add_style("TestStyle", WD_STYLE_TYPE.PARAGRAPH)

    # Add a table with 2 rows, 2 columns
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell 0,0"
    table.cell(0, 1).text = "Cell 0,1"
    table.cell(1, 0).text = "Cell 1,0"
    table.cell(1, 1).text = "Cell 1,1"

    ir = ChangeStyle(
        path="section0body0table0",
        target_style="TestStyle",
        cell_table_indices=((0, 0, 0), (0, 1, 1)),
    )

    apply_change_style(doc, ir)

    # Check that cell paragraphs are styled
    assert table.cell(0, 0).paragraphs[0].style == target_style
    assert table.cell(1, 1).paragraphs[0].style == target_style
    # Other cells should not be styled
    assert table.cell(0, 1).paragraphs[0].style != target_style
    assert table.cell(1, 0).paragraphs[0].style != target_style


def test_apply_change_style_skips_markers() -> None:
    doc = Document()
    _ensure_style(doc, "VerbatimChar", WD_STYLE_TYPE.CHARACTER)
    verbatim = doc.styles["VerbatimChar"]
    target_style = doc.styles.add_style("TestStyle", WD_STYLE_TYPE.PARAGRAPH)

    marker = doc.add_paragraph("irstart:block1")
    for run in marker.runs:
        run.style = verbatim

    content = doc.add_paragraph("Content")

    ir = ChangeStyle(
        path="section0body0",
        target_style="TestStyle",
        paragraph_indices=(0, 1),
    )

    apply_change_style(doc, ir)

    assert is_marker_paragraph(marker)
    assert marker.style != target_style
    assert content.style == target_style


def test_missing_target_style_raises_valueerror() -> None:
    doc = Document()
    ir = ChangeStyle(
        path="section0body0",
        target_style="NonExistentStyle",
        paragraph_indices=(0,),
    )

    with pytest.raises(ValueError, match="NonExistentStyle"):
        apply_change_style(doc, ir)


def test_empty_indices_does_nothing() -> None:
    doc = Document()
    target_style = doc.styles.add_style("TestStyle", WD_STYLE_TYPE.PARAGRAPH)

    p0 = doc.add_paragraph("Paragraph 0")
    original_style = p0.style

    ir = ChangeStyle(
        path="section0body0",
        target_style="TestStyle",
        paragraph_indices=(),
        cell_table_indices=(),
    )

    apply_change_style(doc, ir)

    assert p0.style == original_style


class TestApplyEditStyle:
    """Tests for apply_edit_style function."""

    @pytest.fixture
    def doc(self) -> Document:
        """Return a fresh Document with a custom test style."""
        d = Document()
        d.styles.add_style("TestStyle", WD_STYLE_TYPE.PARAGRAPH)
        return d

    def test_edit_font_name(self, doc: Document) -> None:
        """Edit an existing style's font_name property."""
        ir = EditStyle(style_name="TestStyle", properties={"font_name": "Arial"})
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.font.name == "Arial"

    def test_edit_bold(self, doc: Document) -> None:
        """Edit an existing style's bold property."""
        ir = EditStyle(style_name="TestStyle", properties={"bold": True})
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.font.bold is True

    def test_edit_alignment(self, doc: Document) -> None:
        """Edit an existing style's alignment property."""
        ir = EditStyle(
            style_name="TestStyle", properties={"alignment": "center"}
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_missing_style_raises_valueerror(self, doc: Document) -> None:
        """Raise ValueError when the target style does not exist."""
        ir = EditStyle(style_name="NonExistent", properties={"font_name": "Arial"})

        with pytest.raises(ValueError, match="NonExistent"):
            apply_edit_style(doc, ir)

    def test_edit_font_size(self, doc: Document) -> None:
        """Edit an existing style's font_size property."""
        ir = EditStyle(style_name="TestStyle", properties={"font_size": 14})
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.font.size.pt == 14

    def test_edit_space_before_and_after(self, doc: Document) -> None:
        """Edit space_before and space_after properties."""
        ir = EditStyle(
            style_name="TestStyle",
            properties={"space_before": 12, "space_after": 6},
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.paragraph_format.space_before.pt == 12
        assert style.paragraph_format.space_after.pt == 6

    def test_edit_line_spacing(self, doc: Document) -> None:
        """Edit line_spacing property."""
        ir = EditStyle(
            style_name="TestStyle",
            properties={"line_spacing": LineSpacing(value=1.5, unit="multiplier")},
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.paragraph_format.line_spacing == 1.5

    def test_edit_font_name_east_asia(self, doc: Document) -> None:
        """Edit font_name_east_asia at the XML level."""
        ir = EditStyle(
            style_name="TestStyle", properties={"font_name_east_asia": "宋体"}
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        rpr = style.element.find(qn("w:rPr"))
        assert rpr is not None
        rfonts = rpr.find(qn("w:rFonts"))
        assert rfonts is not None
        east_asia = rfonts.get(qn("w:eastAsia"))
        assert east_asia == "宋体"

    def test_edit_first_line_indent_chars(self, doc: Document) -> None:
        """Edit first_line_indent_chars at the XML level."""
        ir = EditStyle(
            style_name="TestStyle", properties={"first_line_indent_chars": 2}
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        ppr = style.element.find(qn("w:pPr"))
        assert ppr is not None
        ind = ppr.find(qn("w:ind"))
        assert ind is not None
        first_line_chars = ind.get(qn("w:firstLineChars"))
        assert first_line_chars == "200"

    def test_edit_multiple_properties(self, doc: Document) -> None:
        """Edit multiple properties at once."""
        ir = EditStyle(
            style_name="TestStyle",
            properties={
                "font_name": "Calibri",
                "font_size": 11,
                "bold": False,
                "alignment": "right",
            },
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.font.name == "Calibri"
        assert style.font.size.pt == 11
        assert style.font.bold is False
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_edit_justify_alignment(self, doc: Document) -> None:
        """Edit alignment to justify."""
        ir = EditStyle(
            style_name="TestStyle", properties={"alignment": "justify"}
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_edit_left_alignment(self, doc: Document) -> None:
        """Edit alignment to left."""
        ir = EditStyle(
            style_name="TestStyle", properties={"alignment": "left"}
        )
        apply_edit_style(doc, ir)

        style = doc.styles["TestStyle"]
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT


class TestApplyCreateStyle:
    """Test apply_create_style function."""

    def test_creates_new_style(self) -> None:
        """CreateStyle adds a new paragraph style when it does not exist."""
        doc = Document()
        ir = CreateStyle(
            style_name="CustomStyle",
            base_style=None,
            properties={"font_size": 14},
        )
        apply_create_style(doc, ir)

        assert "CustomStyle" in [s.name for s in doc.styles]
        style = doc.styles["CustomStyle"]
        assert style.font.size.pt == 14

    def test_with_base_style(self) -> None:
        """CreateStyle with base_style inherits from the specified style."""
        doc = Document()
        ir = CreateStyle(
            style_name="DerivedStyle",
            base_style="Normal",
            properties={"font_size": 14},
        )
        apply_create_style(doc, ir)

        style = doc.styles["DerivedStyle"]
        assert style.base_style == doc.styles["Normal"]
        assert style.font.size.pt == 14

    def test_applies_font_name(self) -> None:
        """CreateStyle applies font_name property."""
        doc = Document()
        ir = CreateStyle(
            style_name="FontStyle",
            base_style=None,
            properties={"font_name": "Arial"},
        )
        apply_create_style(doc, ir)

        style = doc.styles["FontStyle"]
        assert style.font.name == "Arial"

    def test_applies_font_size(self) -> None:
        """CreateStyle applies font_size property."""
        doc = Document()
        ir = CreateStyle(
            style_name="SizeStyle",
            base_style=None,
            properties={"font_size": 16},
        )
        apply_create_style(doc, ir)

        style = doc.styles["SizeStyle"]
        assert style.font.size.pt == 16

    def test_applies_bold(self) -> None:
        """CreateStyle applies bold property."""
        doc = Document()
        ir = CreateStyle(
            style_name="BoldStyle",
            base_style=None,
            properties={"bold": True},
        )
        apply_create_style(doc, ir)

        style = doc.styles["BoldStyle"]
        assert style.font.bold is True

    def test_applies_alignment_center(self) -> None:
        """CreateStyle applies alignment property (center)."""
        doc = Document()
        ir = CreateStyle(
            style_name="CenterStyle",
            base_style=None,
            properties={"alignment": "center"},
        )
        apply_create_style(doc, ir)

        style = doc.styles["CenterStyle"]
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_applies_space_before(self) -> None:
        """CreateStyle applies space_before property."""
        doc = Document()
        ir = CreateStyle(
            style_name="SpaceStyle",
            base_style=None,
            properties={"space_before": 12},
        )
        apply_create_style(doc, ir)

        style = doc.styles["SpaceStyle"]
        assert style.paragraph_format.space_before.pt == 12

    def test_applies_space_after(self) -> None:
        """CreateStyle applies space_after property."""
        doc = Document()
        ir = CreateStyle(
            style_name="SpaceAfterStyle",
            base_style=None,
            properties={"space_after": 6},
        )
        apply_create_style(doc, ir)

        style = doc.styles["SpaceAfterStyle"]
        assert style.paragraph_format.space_after.pt == 6

    def test_applies_line_spacing(self) -> None:
        """CreateStyle applies line_spacing property."""
        doc = Document()
        ir = CreateStyle(
            style_name="LineSpacingStyle",
            base_style=None,
            properties={"line_spacing": LineSpacing(value=1.5, unit="multiplier")},
        )
        apply_create_style(doc, ir)

        style = doc.styles["LineSpacingStyle"]
        assert style.paragraph_format.line_spacing == 1.5

    def test_applies_east_asia_font_xml(self) -> None:
        """CreateStyle applies font_name_east_asia via XML w:eastAsia attribute."""
        doc = Document()
        ir = CreateStyle(
            style_name="EastAsiaStyle",
            base_style=None,
            properties={"font_name_east_asia": "宋体"},
        )
        apply_create_style(doc, ir)

        style = doc.styles["EastAsiaStyle"]
        rpr = style.element.find(qn("w:rPr"))
        assert rpr is not None
        rfonts = rpr.find(qn("w:rFonts"))
        assert rfonts is not None
        assert rfonts.get(qn("w:eastAsia")) == "宋体"

    def test_applies_first_line_indent_chars_xml(self) -> None:
        """CreateStyle applies first_line_indent_chars via XML w:firstLineChars."""
        doc = Document()
        ir = CreateStyle(
            style_name="IndentStyle",
            base_style=None,
            properties={"first_line_indent_chars": 2},
        )
        apply_create_style(doc, ir)

        style = doc.styles["IndentStyle"]
        ppr = style.element.find(qn("w:pPr"))
        assert ppr is not None
        ind = ppr.find(qn("w:ind"))
        assert ind is not None
        assert ind.get(qn("w:firstLineChars")) == "200"

    def test_does_not_overwrite_existing_style(self) -> None:
        """CreateStyle on existing style preserves it and applies new properties."""
        doc = Document()
        style = doc.styles.add_style("ExistingStyle", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Courier New"

        ir = CreateStyle(
            style_name="ExistingStyle",
            base_style=None,
            properties={"font_size": 10},
        )
        apply_create_style(doc, ir)

        style = doc.styles["ExistingStyle"]
        assert style.font.name == "Courier New"
        assert style.font.size.pt == 10

    def test_multiple_properties_together(self) -> None:
        """CreateStyle applies multiple properties at once."""
        doc = Document()
        ir = CreateStyle(
            style_name="MultiStyle",
            base_style=None,
            properties={
                "font_name": "Times New Roman",
                "font_size": 12,
                "bold": True,
                "alignment": "justify",
            },
        )
        apply_create_style(doc, ir)

        style = doc.styles["MultiStyle"]
        assert style.font.name == "Times New Roman"
        assert style.font.size.pt == 12
        assert style.font.bold is True
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


class TestApplyStyleProperties:
    """Test _apply_style_properties helper directly."""

    def test_font_name(self) -> None:
        """_apply_style_properties sets font name."""
        doc = Document()
        style = doc.styles.add_style("TestFont", WD_STYLE_TYPE.PARAGRAPH)
        props = StyleProperties(font_name="Arial")
        _apply_style_properties(style, props)
        assert style.font.name == "Arial"

    def test_font_size(self) -> None:
        """_apply_style_properties sets font size."""
        doc = Document()
        style = doc.styles.add_style("TestSize", WD_STYLE_TYPE.PARAGRAPH)
        props = StyleProperties(font_size=16)
        _apply_style_properties(style, props)
        assert style.font.size.pt == 16

    def test_bold(self) -> None:
        """_apply_style_properties sets bold."""
        doc = Document()
        style = doc.styles.add_style("TestBold", WD_STYLE_TYPE.PARAGRAPH)
        props = StyleProperties(bold=True)
        _apply_style_properties(style, props)
        assert style.font.bold is True

    def test_alignment_center(self) -> None:
        """_apply_style_properties sets alignment to center."""
        doc = Document()
        style = doc.styles.add_style("TestAlign", WD_STYLE_TYPE.PARAGRAPH)
        props = StyleProperties(alignment="center")
        _apply_style_properties(style, props)
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_east_asia_font_xml(self) -> None:
        """_apply_style_properties sets w:eastAsia via XML."""
        doc = Document()
        style = doc.styles.add_style("TestEastAsia", WD_STYLE_TYPE.PARAGRAPH)
        props = StyleProperties(font_name_east_asia="宋体")
        _apply_style_properties(style, props)

        rpr = style.element.find(qn("w:rPr"))
        assert rpr is not None
        rfonts = rpr.find(qn("w:rFonts"))
        assert rfonts is not None
        assert rfonts.get(qn("w:eastAsia")) == "宋体"

    def test_first_line_indent_chars_xml(self) -> None:
        """_apply_style_properties sets w:firstLineChars via XML."""
        doc = Document()
        style = doc.styles.add_style("TestIndent", WD_STYLE_TYPE.PARAGRAPH)
        props = StyleProperties(first_line_indent_chars=2)
        _apply_style_properties(style, props)

        ppr = style.element.find(qn("w:pPr"))
        assert ppr is not None
        ind = ppr.find(qn("w:ind"))
        assert ind is not None
        assert ind.get(qn("w:firstLineChars")) == "200"


def _get_tbl_borders(table):
    tbl_pr = table._tbl.tblPr
    return tbl_pr.find(qn("w:tblBorders"))


def _get_tbl_alignment(table):
    tbl_pr = table._tbl.tblPr
    jc = tbl_pr.find(qn("w:jc"))
    return jc.get(qn("w:val")) if jc is not None else None


class TestApplyTableFormat:
    """Tests for apply_table_format function."""

    def test_adds_tbl_borders_with_six_edges(self) -> None:
        """apply_table_format adds w:tblBorders with 6 child elements."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(),
        )
        apply_table_format(doc, ir)

        borders = _get_tbl_borders(table)
        assert borders is not None
        assert len(borders) == 6

    def test_border_edges_have_correct_attributes(self) -> None:
        """Each border edge has correct attributes."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(),
        )
        apply_table_format(doc, ir)

        borders = _get_tbl_borders(table)
        assert borders is not None
        for edge in borders:
            assert edge.get(qn("w:val")) == "single"
            assert edge.get(qn("w:sz")) == "4"
            assert edge.get(qn("w:color")) == "000000"
            assert edge.get(qn("w:space")) == "0"

    def test_adds_center_alignment(self) -> None:
        """w:jc with w:val='center' is set in w:tblPr."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(),
        )
        apply_table_format(doc, ir)

        assert _get_tbl_alignment(table) == "center"

    def test_creates_tbl_pr_when_missing(self) -> None:
        """A table without w:tblPr gets one created."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        tbl = table._tbl
        if tbl.tblPr is not None:
            tbl.remove(tbl.tblPr)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(),
        )
        apply_table_format(doc, ir)

        assert tbl.tblPr is not None
        borders = _get_tbl_borders(table)
        assert borders is not None

    def test_overwrites_existing_tbl_borders(self) -> None:
        """Existing w:tblBorders is overwritten, not duplicated."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        tbl_pr = table._tbl.tblPr
        existing = OxmlElement("w:tblBorders")
        existing.append(OxmlElement("w:top"))
        tbl_pr.append(existing)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(),
        )
        apply_table_format(doc, ir)

        borders = _get_tbl_borders(table)
        assert borders is not None
        assert len(borders) == 6

    def test_alignment_none_does_not_add_jc(self) -> None:
        """alignment=None means no w:jc element added."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(alignment=None),
        )
        apply_table_format(doc, ir)

        tbl_pr = table._tbl.tblPr
        jc = tbl_pr.find(qn("w:jc"))
        assert jc is None

    def test_borders_none_does_not_add_tbl_borders(self) -> None:
        """borders=None means no w:tblBorders element added."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)

        ir = ApplyTableFormat(
            path="t0",
            table_indices=(0,),
            format=TableFormat(borders=None),
        )
        apply_table_format(doc, ir)

        borders = _get_tbl_borders(table)
        assert borders is None
