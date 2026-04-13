from __future__ import annotations

import logging
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document

from typst_ir.mapping.types import DocxTypstMapping

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)


def _extract_all_paragraphs(docx_path: Path) -> list[Paragraph]:
    doc = Document(str(docx_path))
    paragraphs: list[Paragraph] = []

    for paragraph in doc.paragraphs:
        paragraphs.append(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraphs.append(paragraph)

    return paragraphs


def locate_nodes(mapping: DocxTypstMapping, docx_path: Path) -> dict[str, int]:
    """Map each NodeID to its start paragraph index in the DOCX document.

    Args:
        mapping: DOCX-to-Typst IR mapping containing node entries.
        docx_path: Path to the source ``.docx`` file.

    Returns:
        Dictionary mapping NodeID canonical string to ``source_start`` index.

    Raises:
        FileNotFoundError: If ``docx_path`` does not exist.
        ValueError: If ``docx_path`` is not a valid DOCX file.
    """
    if not docx_path.exists():
        raise FileNotFoundError(f"Input file not found: {docx_path}")

    try:
        with zipfile.ZipFile(docx_path, "r") as zf:
            if "word/document.xml" not in zf.namelist():
                raise ValueError(f"Invalid DOCX file: {docx_path}")
    except zipfile.BadZipFile as exc:
        raise ValueError(f"Invalid DOCX file: {docx_path}") from exc

    doc_paragraphs = _extract_all_paragraphs(docx_path)
    doc_length = len(doc_paragraphs)

    result: dict[str, int] = {}
    for entry in mapping.entries.values():
        node_id_str = entry.node_id.canonical
        result[node_id_str] = entry.source_start
        if entry.source_end > doc_length:
            logger.warning(
                "Mapping entry %s has source_end=%d which exceeds "
                "document paragraph count=%d",
                node_id_str,
                entry.source_end,
                doc_length,
            )

    return result


def validate_mapping(mapping: DocxTypstMapping, docx_path: Path) -> None:
    """Validate that mapping paragraph indices are within document bounds.

    Args:
        mapping: DOCX-to-Typst IR mapping to validate.
        docx_path: Path to the source ``.docx`` file.

    Raises:
        FileNotFoundError: If ``docx_path`` does not exist.
        ValueError: If ``docx_path`` is not a valid DOCX file, or if any
            mapping entry references a paragraph index outside the document.
    """
    if not docx_path.exists():
        raise FileNotFoundError(f"Input file not found: {docx_path}")

    try:
        with zipfile.ZipFile(docx_path, "r") as zf:
            if "word/document.xml" not in zf.namelist():
                raise ValueError(f"Invalid DOCX file: {docx_path}")
    except zipfile.BadZipFile as exc:
        raise ValueError(f"Invalid DOCX file: {docx_path}") from exc

    doc_paragraphs = _extract_all_paragraphs(docx_path)
    doc_length = len(doc_paragraphs)

    for entry in mapping.entries.values():
        if entry.source_start < 0 or entry.source_end < 0:
            raise ValueError(
                f"Mapping entry {entry.node_id.canonical} has negative index: "
                f"source_start={entry.source_start}, source_end={entry.source_end}"
            )
        if entry.source_start > doc_length:
            raise ValueError(
                f"Mapping entry {entry.node_id.canonical} has source_start="
                f"{entry.source_start} which exceeds document paragraph "
                f"count={doc_length}"
            )
        if entry.source_end > doc_length:
            raise ValueError(
                f"Mapping entry {entry.node_id.canonical} has source_end="
                f"{entry.source_end} which exceeds document paragraph "
                f"count={doc_length}"
            )
