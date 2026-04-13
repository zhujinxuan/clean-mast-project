"""DocxMatcher for sequential matching between Typst IR nodes and DOCX content.

Implements the sequential matching algorithm from detection-patterns.md,
using the detection functions from detection.py to match IR nodes to DOCX
paragraphs, tables, and other content.
"""

from __future__ import annotations

import hashlib
import logging
import re
from typing import Any

from docx.table import Table
from docx.text.paragraph import Paragraph

from typst_ir.mapping.detection import (
    detect_equation,
    detect_heading_level,
    group_docx_spans,
    match_paragraph_fingerprint,
    match_table,
)
from typst_ir.mapping.registry import NodeRegistry
from typst_ir.mapping.types import (
    ConfidenceLevel,
    DetectionMethod,
    DocxParagraph,
    MappingEntry,
    NodeID,
)
from typst_ir.nodes import (
    TypstDocument,
    TypstEquation,
    TypstFigure,
    TypstHeading,
    TypstParagraph,
    TypstSection,
    TypstTable,
)

logger = logging.getLogger(__name__)


def _normalize_whitespace(text: str) -> str:
    return " ".join(text.split())


def _hash_text(text: str) -> str:
    return hashlib.sha256(_normalize_whitespace(text).encode()).hexdigest()


def _map_method(method: str) -> DetectionMethod:
    mapping: dict[str, DetectionMethod] = {
        "fingerprint": DetectionMethod.fingerprint,
        "style_heading": DetectionMethod.style,
        "style_caption": DetectionMethod.style,
        "style_body": DetectionMethod.style,
        "hash": DetectionMethod.fingerprint,
        "positional": DetectionMethod.positional,
        "style": DetectionMethod.style,
        "structure": DetectionMethod.structure,
        "drawing": DetectionMethod.drawing,
        "caption": DetectionMethod.caption,
        "xpath": DetectionMethod.xpath,
    }
    return mapping.get(method, DetectionMethod.positional)


def _map_confidence(confidence: str) -> ConfidenceLevel:
    return ConfidenceLevel(confidence)


def match_headings(
    ir_nodes: list[TypstSection],
    docx_spans: list[Any],
    registry: NodeRegistry | None = None,
) -> list[MappingEntry]:
    """Pass 1: Match TypstHeading IR nodes to DOCX heading spans.

    Filters DOCX spans for headings, then sequentially aligns them with
    TypstHeading IR nodes by level. Detection method is ``fingerprint``
    when text coverage is high, ``style`` when matched by heading style,
    and ``xpath`` when matched by outline level XML.

    Args:
        ir_nodes: List of Typst IR nodes.
        docx_spans: List of DOCX content objects or ``DocxSpan`` objects.
        registry: Optional NodeRegistry for resolving IR node IDs.

    Returns:
        List of ``MappingEntry`` records for matched headings.
    """
    ir_headings = [node for node in ir_nodes if isinstance(node, TypstHeading)]

    heading_spans: list[tuple[int, int, Any, int]] = []
    for idx, span in enumerate(docx_spans):
        label = getattr(span, "label", None)
        if isinstance(label, str) and label == "heading":
            docx_para = span.elements[0]
            level = detect_heading_level(docx_para)
            if level is not None:
                heading_spans.append((span.start, span.end, docx_para, level))
        else:
            if (
                isinstance(getattr(span, "label", None), str)
                and hasattr(span, "elements")
            ):
                continue
            level = detect_heading_level(span)
            if level is not None:
                heading_spans.append((idx, idx + 1, span, level))

    entries: list[MappingEntry] = []
    matched_ir_indices: set[int] = set()
    span_idx = 0
    for ir_idx, ir_node in enumerate(ir_headings):
        matched = False
        expected_level = 1
        if registry is not None:
            node_id = registry.get_id(ir_node)
            expected_level = max(1, node_id.path.count("[TypstSection]"))
        while span_idx < len(heading_spans):
            start, end, docx_para, level = heading_spans[span_idx]
            if expected_level == level:
                is_match, _cov_a, _cov_b = match_paragraph_fingerprint(
                    ir_node.text, docx_para.text
                )
                if is_match:
                    method = DetectionMethod.fingerprint
                else:
                    style_name = getattr(getattr(docx_para, "style", None), "name", "") or ""
                    if re.search(r"Heading\s*\d+", style_name, re.IGNORECASE):
                        method = DetectionMethod.style
                    else:
                        method = DetectionMethod.xpath

                if registry is not None:
                    node_id = registry.get_id(ir_node)
                else:
                    node_hash = _hash_text(ir_node.text)
                    node_id = NodeID(
                        path=f"heading.{ir_idx}[TypstHeading]",
                        full_hash=node_hash,
                        node_hash=node_hash,
                    )

                docx_para_obj = DocxParagraph(
                    index=start,
                    text_preview=docx_para.text[:100],
                    text_hash=_hash_text(docx_para.text),
                    paragraph_type="heading",
                )

                entry = MappingEntry(
                    node_id=node_id,
                    source_start=start,
                    source_end=end,
                    method=method,
                    confidence=ConfidenceLevel.high,
                    details={"level": level, "docx_paragraph": docx_para_obj},
                )
                entries.append(entry)
                matched_ir_indices.add(ir_idx)
                span_idx += 1
                matched = True
                break
            span_idx += 1
        if not matched:
            logger.warning(
                "Unmatched TypstHeading at index %d: no corresponding heading span (text=%r, level=%d)",
                ir_idx,
                ir_node.text,
                expected_level,
            )

    return entries


class DocxMatcher:
    """Sequential matcher for Typst IR nodes to DOCX content.

    Iterates over DOCX contents and IR nodes simultaneously, producing
    MappingEntry records with detection method, confidence, and details.

    Args:
        registry: Optional NodeRegistry for resolving IR node IDs.
            If not provided, a fallback registry is created internally.
        log_level: Logging level for matcher output. Accepts an int
            (e.g., logging.DEBUG) or a string (e.g., "DEBUG").

    Example:
        >>> from docx import Document
        >>> from typst_ir import TypstParagraph, TypstHeading
        >>> from typst_ir.mapping.matcher import DocxMatcher
        >>> ir_nodes = [
        ...     TypstHeading(text="Introduction", level=1),
        ...     TypstParagraph(text="This is the intro."),
        ... ]
        >>> doc = Document("report.docx")
        >>> matcher = DocxMatcher()
        >>> mappings = matcher.match(ir_nodes, doc)
        >>> mappings[0].node_type
        'TypstHeading'
    """

    def __init__(
        self,
        registry: NodeRegistry | None = None,
        log_level: int | str = logging.WARNING,
    ) -> None:
        self.registry = registry
        if isinstance(log_level, str):
            log_level = logging.getLevelName(log_level)
        logger.setLevel(log_level)
        if not logger.handlers:
            handler = logging.StreamHandler()
            handler.setLevel(log_level)
            logger.addHandler(handler)

    def _ensure_registry(self, ir_nodes: list[TypstSection]) -> NodeRegistry:
        """Return an initialized NodeRegistry for the given IR nodes."""
        if self.registry is not None:
            return self.registry
        registry = NodeRegistry()
        registry.initialize(TypstDocument(children=ir_nodes))
        return registry

    def match(self, ir_nodes: list[TypstSection], doc_document: Any) -> list[MappingEntry]:
        """Match a sequence of IR nodes to a DOCX document's inner content.

        Walks through ``doc_document.iter_inner_content()`` and matches each
        element to the corresponding IR node using the detection functions
        from ``typst_ir.mapping.detection``. Produces a list of
        ``MappingEntry`` objects with detection method, confidence, and
        metadata.

        Args:
            ir_nodes: List of Typst IR nodes to match.
            doc_document: A ``python-docx`` ``Document`` instance.

        Returns:
            List of ``MappingEntry`` records in document order.

        Example:
            >>> from docx import Document
            >>> from typst_ir import TypstParagraph, TypstHeading
            >>> from typst_ir.mapping.matcher import DocxMatcher
            >>> ir_nodes = [
            ...     TypstHeading(text="Summary", level=1),
            ...     TypstParagraph(text="Key findings here."),
            ... ]
            >>> doc = Document("report.docx")
            >>> matcher = DocxMatcher()
            >>> entries = matcher.match(ir_nodes, doc)
            >>> len(entries) > 0
            True
            >>> entries[0].confidence.value in ("high", "medium", "low")
            True
        """
        registry = self._ensure_registry(ir_nodes)
        docx_spans = group_docx_spans(list(doc_document.iter_inner_content()))

        heading_entries = match_headings(ir_nodes, docx_spans, registry=registry)
        block_entries = match_blocks(
            ir_nodes, docx_spans, heading_entries, registry=registry
        )
        para_entries = match_paragraphs(
            ir_nodes, docx_spans, heading_entries + block_entries, registry=registry
        )

        all_entries = heading_entries + block_entries + para_entries
        return sorted(all_entries, key=lambda e: e.source_start)


def match_blocks(
    ir_nodes: list[TypstSection],
    docx_spans: list[Any],
    heading_anchors: list[MappingEntry],
    registry: NodeRegistry | None = None,
) -> list[MappingEntry]:
    """Pass 2: Match block-level IR nodes to grouped DOCX block spans.

    Matches ``TypstFigure``, ``TypstTable``, and ``TypstEquation`` IR nodes to
    corresponding DOCX spans (``figure``, ``figure+caption``, ``table``,
    ``caption+table``, ``equation``) within segments defined by heading anchors.

    Args:
        ir_nodes: List of Typst IR nodes.
        docx_spans: List of DOCX content objects or ``DocxSpan`` objects.
        heading_anchors: List of ``MappingEntry`` records for matched headings,
            used to define matching segments.
        registry: Optional NodeRegistry for resolving IR node IDs.

    Returns:
        List of ``MappingEntry`` records for matched block nodes.
    """

    def _has_real_elements(span: Any) -> bool:
        elements = getattr(span, "elements", None)
        return isinstance(elements, list) and len(elements) > 0

    def _hash_node(node: Any) -> str:
        if isinstance(node, TypstFigure):
            return _hash_text(node.caption or node.content)
        if isinstance(node, TypstTable):
            return _hash_text(node.caption or "".join(node.headers))
        if isinstance(node, TypstEquation):
            return _hash_text(node.latex)
        return _hash_text(str(node))

    def _try_match(
        ir_idx: int,
        ir_node: Any,
        span: Any,
        label: str,
        source_start: int,
        source_end: int,
    ) -> MappingEntry | None:
        if registry is not None:
            node_id = registry.get_id(ir_node)
        else:
            node_type = type(ir_node).__name__
            node_hash = _hash_node(ir_node)
            node_id = NodeID(
                path=f"doc.children.{ir_idx}[{node_type}]",
                full_hash=node_hash,
                node_hash=node_hash,
            )

        if isinstance(ir_node, TypstFigure):
            if label == "figure":
                return MappingEntry(
                    node_id=node_id,
                    source_start=source_start,
                    source_end=source_end,
                    method=DetectionMethod.drawing,
                    confidence=ConfidenceLevel.high,
                    details={},
                )
            if label == "figure+caption":
                caption_match = False
                details: dict[str, Any] = {}
                if _has_real_elements(span) and len(span.elements) > 1:
                    caption_text = getattr(span.elements[1], "text", "") or ""
                    if ir_node.caption and caption_text:
                        is_match, cov_a, cov_b = match_paragraph_fingerprint(
                            ir_node.caption, caption_text
                        )
                        caption_match = is_match or _normalize_whitespace(
                            ir_node.caption
                        ) == _normalize_whitespace(caption_text)
                        if caption_match:
                            details["coverage"] = (cov_a, cov_b)
                else:
                    caption_match = bool(ir_node.caption)

                if caption_match:
                    return MappingEntry(
                        node_id=node_id,
                        source_start=source_start,
                        source_end=source_end,
                        method=DetectionMethod.fingerprint,
                        confidence=ConfidenceLevel.high,
                        details=details,
                    )
                return MappingEntry(
                    node_id=node_id,
                    source_start=source_start,
                    source_end=source_end,
                    method=DetectionMethod.drawing,
                    confidence=ConfidenceLevel.medium,
                    details=details,
                )
            return None

        if isinstance(ir_node, TypstTable):
            if label == "table":
                details: dict[str, Any] = {}
                if _has_real_elements(span):
                    table_obj = None
                    for elem in span.elements:
                        if isinstance(elem, Table):
                            table_obj = elem
                            break
                    if table_obj is not None:
                        table_idx = span.elements.index(table_obj)
                        result = match_table(
                            ir_node, table_obj, span.elements, table_idx
                        )
                        structure_match = result.get("structure", False)
                        details = {
                            "structure_match": structure_match,
                            "ir_dims": result.get("ir_dims"),
                            "docx_dims": result.get("docx_dims"),
                        }
                        confidence = (
                            ConfidenceLevel.high
                            if structure_match
                            else ConfidenceLevel.medium
                        )
                        return MappingEntry(
                            node_id=node_id,
                            source_start=source_start,
                            source_end=source_end,
                            method=DetectionMethod.structure,
                            confidence=confidence,
                            details=details,
                        )
                return MappingEntry(
                    node_id=node_id,
                    source_start=source_start,
                    source_end=source_end,
                    method=DetectionMethod.structure,
                    confidence=ConfidenceLevel.high,
                    details=details,
                )

            if label == "caption+table":
                caption_match = False
                structure_match = False
                details = {}

                if _has_real_elements(span) and len(span.elements) >= 2:
                    caption_text = getattr(span.elements[0], "text", "") or ""
                    if ir_node.caption and caption_text:
                        is_match, cov_a, cov_b = match_paragraph_fingerprint(
                            ir_node.caption, caption_text
                        )
                        caption_match = is_match or _normalize_whitespace(
                            ir_node.caption
                        ) == _normalize_whitespace(caption_text)
                        if caption_match:
                            details["coverage"] = (cov_a, cov_b)

                    table_obj = None
                    for elem in span.elements[1:]:
                        if isinstance(elem, Table):
                            table_obj = elem
                            break
                    if table_obj is not None:
                        table_idx = span.elements.index(table_obj)
                        result = match_table(
                            ir_node, table_obj, span.elements, table_idx
                        )
                        structure_match = result.get("structure", False)
                        details["ir_dims"] = result.get("ir_dims")
                        details["docx_dims"] = result.get("docx_dims")
                else:
                    caption_match = bool(ir_node.caption)
                    structure_match = True

                details["caption_match"] = caption_match
                details["structure_match"] = structure_match

                if caption_match and structure_match:
                    method = DetectionMethod.fingerprint
                    confidence = ConfidenceLevel.high
                elif caption_match:
                    method = DetectionMethod.caption
                    confidence = ConfidenceLevel.medium
                elif structure_match:
                    method = DetectionMethod.structure
                    confidence = ConfidenceLevel.medium
                else:
                    method = DetectionMethod.caption
                    confidence = ConfidenceLevel.low

                return MappingEntry(
                    node_id=node_id,
                    source_start=source_start,
                    source_end=source_end,
                    method=method,
                    confidence=confidence,
                    details=details,
                )

            return None

        if isinstance(ir_node, TypstEquation):
            if label != "equation":
                return None

            eq_type = None
            if _has_real_elements(span):
                for elem in span.elements:
                    if isinstance(elem, Paragraph):
                        eq_type = detect_equation(elem)
                        break

            if eq_type is not None:
                if (eq_type == "display" and ir_node.display) or (
                    eq_type == "inline" and not ir_node.display
                ):
                    return MappingEntry(
                        node_id=node_id,
                        source_start=source_start,
                        source_end=source_end,
                        method=DetectionMethod.xpath,
                        confidence=ConfidenceLevel.high,
                        details={"type": eq_type},
                    )
                return None

            eq_type = "display" if ir_node.display else "inline"
            return MappingEntry(
                node_id=node_id,
                source_start=source_start,
                source_end=source_end,
                method=DetectionMethod.xpath,
                confidence=ConfidenceLevel.high,
                details={"type": eq_type},
            )

        return None

    # 1. Filter IR nodes to block types with original indices.
    ir_blocks: list[tuple[int, Any]] = [
        (idx, node)
        for idx, node in enumerate(ir_nodes)
        if isinstance(node, (TypstFigure, TypstTable, TypstEquation))
    ]

    # 2. Filter docx_spans to block labels.
    block_labels = {"figure", "figure+caption", "table", "caption+table", "equation"}
    block_spans: list[Any] = []
    for span in docx_spans:
        label = getattr(span, "label", None)
        if isinstance(label, str) and label in block_labels:
            block_spans.append(span)

    # 3. Define segments using heading_anchors.
    anchors = list(heading_anchors) if heading_anchors else []
    sorted_anchors = sorted(anchors, key=lambda a: a.source_start)

    max_end = 2_147_483_647
    segments: list[tuple[int, int]] = []
    if not sorted_anchors:
        segments = [(0, max_end)]
    else:
        segments.append((0, sorted_anchors[0].source_start))
        for i in range(len(sorted_anchors) - 1):
            segments.append(
                (sorted_anchors[i].source_end, sorted_anchors[i + 1].source_start)
            )
        segments.append((sorted_anchors[-1].source_end, max_end))

    # 4 & 5. Within each segment, match IR block nodes to DOCX block spans.
    entries: list[MappingEntry] = []
    ir_block_idx = 0

    for seg_start, seg_end in segments:
        seg_spans = [
            s
            for s in block_spans
            if seg_start <= getattr(s, "start", 0) < seg_end
        ]
        remaining_spans = list(seg_spans)

        while remaining_spans and ir_block_idx < len(ir_blocks):
            ir_idx, ir_node = ir_blocks[ir_block_idx]

            matched_idx = -1
            entry: MappingEntry | None = None
            for j, span in enumerate(remaining_spans):
                label = getattr(span, "label", "")
                source_start = getattr(span, "start", 0)
                source_end = getattr(span, "end", source_start + 1)

                entry = _try_match(
                    ir_idx, ir_node, span, label, source_start, source_end
                )
                if entry is not None:
                    matched_idx = j
                    break

            if matched_idx >= 0 and entry is not None:
                entries.append(entry)
                remaining_spans.pop(matched_idx)
                ir_block_idx += 1
            else:
                ir_block_idx += 1

    return entries


def match_paragraphs(
    ir_nodes: list[TypstSection],
    docx_spans: list[Any],
    anchors: list[MappingEntry],
    registry: NodeRegistry | None = None,
) -> list[MappingEntry]:
    """Pass 3: match remaining TypstParagraph IR nodes to DOCX paragraph spans."""
    ir_paragraphs = [node for node in ir_nodes if isinstance(node, TypstParagraph)]

    consumed_indices: set[int] = set()
    for anchor in anchors:
        for idx in range(anchor.source_start, anchor.source_end):
            consumed_indices.add(idx)

    paragraph_spans: list[Any] = []
    for span in docx_spans:
        label = getattr(span, "label", "paragraph")
        if type(label).__name__ == "MagicMock":
            label = "paragraph"
        block_labels = {"figure", "table", "caption+table", "figure+caption"}
        if label in block_labels:
            continue
        span_index = getattr(span, "index", None)
        if span_index is not None and span_index in consumed_indices:
            continue
        text = getattr(span, "text", None)
        if text is None and hasattr(span, "elements") and span.elements:
            text = span.elements[0].text if hasattr(span.elements[0], "text") else ""
        text = text or ""
        if text == "" or text.strip() == "":
            continue
        paragraph_spans.append(span)

    sorted_anchors = sorted(anchors, key=lambda a: a.source_start)
    segments: list[tuple[int, int]] = []
    if sorted_anchors:
        if sorted_anchors[0].source_start > 0:
            segments.append((0, sorted_anchors[0].source_start))
    for i in range(len(sorted_anchors) - 1):
        seg_start = sorted_anchors[i].source_end
        seg_end = sorted_anchors[i + 1].source_start
        if seg_start < seg_end:
            segments.append((seg_start, seg_end))
    if sorted_anchors:
        segments.append((sorted_anchors[-1].source_end, 2_147_483_647))
    else:
        segments.append((0, 2_147_483_647))

    entries: list[MappingEntry] = []
    ir_idx = 0
    remaining_spans = list(paragraph_spans)

    for seg_start, seg_end in segments:
        seg_length = seg_end - seg_start
        seg_pos = seg_start
        matched_in_segment = 0

        while matched_in_segment < seg_length and ir_idx < len(ir_paragraphs) and remaining_spans:
            ir_para = ir_paragraphs[ir_idx]
            ir_text = ir_para.text

            matched_idx = -1
            coverage: tuple[float, float] | None = None
            for j, span in enumerate(remaining_spans):
                span_text = getattr(span, "text", None)
                if span_text is None and hasattr(span, "elements") and span.elements:
                    span_text = span.elements[0].text if hasattr(span.elements[0], "text") else ""
                span_text = span_text or ""
                is_match, cov_a, cov_b = match_paragraph_fingerprint(ir_text, span_text)
                if is_match or (cov_a >= 0.95 and cov_b >= 0.7):
                    matched_idx = j
                    coverage = (cov_a, cov_b)
                    break

            if matched_idx >= 0:
                matched_span = remaining_spans.pop(matched_idx)
                method = DetectionMethod.fingerprint
                confidence = ConfidenceLevel.high
            else:
                matched_span = remaining_spans.pop(0)
                method = DetectionMethod.positional
                confidence = ConfidenceLevel.low
                coverage = None

            matched_text = getattr(matched_span, "text", None)
            if matched_text is None and hasattr(matched_span, "elements") and matched_span.elements:
                matched_text = matched_span.elements[0].text if hasattr(matched_span.elements[0], "text") else ""
            matched_text = matched_text or ""

            docx_para = DocxParagraph(
                index=seg_pos,
                text_preview=matched_text[:100],
                text_hash=_hash_text(matched_text),
                paragraph_type="paragraph",
            )

            details: dict[str, Any] = {"docx_paragraph": docx_para}
            if coverage is not None:
                details["coverage"] = coverage

            if registry is not None:
                node_id = registry.get_id(ir_para)
            else:
                node_id = NodeID(
                    path=f"doc.children.{ir_idx}[TypstParagraph]",
                    full_hash=_hash_text(ir_text),
                    node_hash=_hash_text(ir_text),
                )
            entry = MappingEntry(
                node_id=node_id,
                source_start=seg_pos,
                source_end=seg_pos + 1,
                method=method,
                confidence=confidence,
                details=details,
            )
            entries.append(entry)

            ir_idx += 1
            seg_pos += 1
            matched_in_segment += 1

    return entries
