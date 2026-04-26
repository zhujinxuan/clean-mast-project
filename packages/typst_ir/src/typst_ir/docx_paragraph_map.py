"""DOCX paragraph range extraction from IR markers.

Provides ``DocxBlockRange`` and ``extract_paragraph_ranges`` for scanning
a DOCX file and extracting the paragraph index ranges between IR markers.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from docx.document import Document

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class DocxBlockRange:
    """Range of paragraph indices for a single block in a DOCX.

    Attributes:
        path: IR path identifier (e.g. ``section0body2table0``).
        begin: Index of the first paragraph inside the block (after start marker).
        end: Index of the end marker paragraph (exclusive of content).
        table_indices: Indices into ``document.tables`` for tables within range.
    """

    path: str
    begin: int
    end: int
    table_indices: tuple[int, ...] = ()


def _find_tables_in_range(doc: Document, start_marker_idx: int, end_marker_idx: int) -> tuple[int, ...]:
    """Return table indices between start and end marker paragraphs."""
    para_elements = {p._element: i for i, p in enumerate(doc.paragraphs)}
    tbl_elements = {t._element: i for i, t in enumerate(doc.tables)}

    in_range = False
    tables: list[int] = []

    for child in doc.element.body:
        tag = getattr(child, "tag", "")
        if tag.endswith("p") and child in para_elements:
            para_idx = para_elements[child]
            if para_idx == start_marker_idx:
                in_range = True
            elif para_idx == end_marker_idx:
                in_range = False
        elif tag.endswith("tbl") and child in tbl_elements:
            if in_range:
                tables.append(tbl_elements[child])

    return tuple(tables)


def extract_paragraph_ranges(docx_path: Path) -> list[DocxBlockRange]:
    """Scan a DOCX for IR markers and extract paragraph ranges.

    Args:
        docx_path: Path to the DOCX file to scan.

    Returns:
        List of ``DocxBlockRange`` for each matched block.
    """
    from docx import Document

    from typst_ir.marker_detection import is_marker_paragraph, parse_marker

    doc = Document(str(docx_path))

    markers: list[tuple[int, str, str]] = []
    for para_idx, para in enumerate(doc.paragraphs):
        if is_marker_paragraph(para):
            parsed = parse_marker(para.runs[0].text)
            if parsed is not None:
                kind, block_id = parsed
                markers.append((para_idx, kind, block_id))

    ranges: list[DocxBlockRange] = []
    stack: list[tuple[int, str]] = []

    for para_idx, kind, block_id in markers:
        if kind == "irstart":
            stack.append((para_idx, block_id))
        elif kind == "irend":
            matched = False
            for i in range(len(stack) - 1, -1, -1):
                if stack[i][1] == block_id:
                    start_idx, _ = stack.pop(i)
                    tables = _find_tables_in_range(doc, start_idx, para_idx)
                    ranges.append(
                        DocxBlockRange(
                            path=block_id,
                            begin=start_idx + 1,
                            end=para_idx,
                            table_indices=tables,
                        )
                    )
                    matched = True
                    break
            if not matched:
                logger.warning(
                    "Unmatched irend marker for block '%s' at paragraph %d",
                    block_id,
                    para_idx,
                )

    for start_idx, block_id in stack:
        logger.warning(
            "Unmatched irstart marker for block '%s' at paragraph %d",
            block_id,
            start_idx,
        )

    return ranges


__all__ = [
    "DocxBlockRange",
    "extract_paragraph_ranges",
]
