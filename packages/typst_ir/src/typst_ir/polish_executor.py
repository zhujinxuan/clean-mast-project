from __future__ import annotations

from typing import Any

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .marker_detection import is_marker_paragraph
from .polish_ir import ApplyTableFormat, ChangeStyle, CreateStyle, EditStyle
from .polish_styles import StyleProperties

_ALIGNMENT_MAP: dict[str, WD_ALIGN_PARAGRAPH] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _ensure_child(parent, tag: str):
    """Find or create a child element with the given tag under parent."""
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _apply_style_properties(style, properties: StyleProperties) -> None:
    """Apply ``StyleProperties`` to a python-docx style object.

    Handles both high-level python-docx API properties and XML-level
    properties that are not exposed by the library (eastAsia font,
    firstLineChars indent).
    """
    if properties.font_name is not None:
        style.font.name = properties.font_name

    if properties.font_size is not None:
        style.font.size = Pt(properties.font_size)

    if properties.bold is not None:
        style.font.bold = properties.bold

    if properties.font_color is not None:
        color_str = properties.font_color.lower()
        if color_str == "black":
            style.font.color.rgb = RGBColor(0, 0, 0)
        elif color_str.startswith("#") and len(color_str) == 7:
            style.font.color.rgb = RGBColor.from_string(color_str[1:])

    if properties.alignment is not None:
        style.paragraph_format.alignment = _ALIGNMENT_MAP[properties.alignment]

    if properties.space_before is not None:
        style.paragraph_format.space_before = Pt(properties.space_before)

    if properties.space_after is not None:
        style.paragraph_format.space_after = Pt(properties.space_after)

    if properties.line_spacing is not None:
        ls = properties.line_spacing
        if ls.unit == "pt":
            style.paragraph_format.line_spacing = Pt(ls.value)
        else:
            style.paragraph_format.line_spacing = ls.value

    # XML-level: East Asian font name (w:eastAsia on w:rPr/w:rFonts)
    if properties.font_name_east_asia is not None:
        rpr = _ensure_child(style.element, "w:rPr")
        rfonts = _ensure_child(rpr, "w:rFonts")
        rfonts.set(qn("w:eastAsia"), properties.font_name_east_asia)

    # XML-level: first-line indent in characters (w:firstLineChars on w:pPr/w:ind)
    if properties.first_line_indent_chars is not None:
        ppr = _ensure_child(style.element, "w:pPr")
        ind = _ensure_child(ppr, "w:ind")
        ind.set(qn("w:firstLineChars"), str(properties.first_line_indent_chars * 100))


def apply_create_style(doc: Any, ir: CreateStyle) -> None:
    """Apply a ``CreateStyle`` IR operation to a DOCX document.

    Creates the paragraph style if it does not already exist, optionally
    inherits from a base style, then applies all specified properties.
    """
    # Convert dict properties to StyleProperties if needed
    if isinstance(ir.properties, dict):
        properties = StyleProperties(**ir.properties)
    else:
        properties = ir.properties

    # Create style if missing
    try:
        style = doc.styles[ir.style_name]
    except KeyError:
        style = doc.styles.add_style(ir.style_name, WD_STYLE_TYPE.PARAGRAPH)

    # Apply base style if provided
    if ir.base_style is not None:
        style.base_style = doc.styles[ir.base_style]

    _apply_style_properties(style, properties)


def apply_change_style(doc: Any, ir: ChangeStyle) -> None:
    """Apply a ``ChangeStyle`` IR operation to a DOCX document.

    Styles paragraphs at ``paragraph_indices`` and table cells at
    ``cell_table_indices`` with the target style.

    Args:
        doc: The python-docx ``Document`` to mutate.
        ir: The ``ChangeStyle`` instruction.

    Raises:
        ValueError: If *ir.target_style* does not exist in *doc.styles*.
    """
    try:
        target_style = doc.styles[ir.target_style]
    except KeyError:
        raise ValueError(
            f"Target style '{ir.target_style}' not found in document styles."
        ) from None

    for idx in ir.paragraph_indices:
        paragraph = doc.paragraphs[idx]
        if is_marker_paragraph(paragraph):
            continue
        paragraph.style = target_style

    for table_idx, row_idx, cell_idx in ir.cell_table_indices:
        cell = doc.tables[table_idx].rows[row_idx].cells[cell_idx]
        for paragraph in cell.paragraphs:
            if is_marker_paragraph(paragraph):
                continue
            paragraph.style = target_style


def apply_edit_style(doc: Any, ir: EditStyle) -> None:
    """Apply an ``EditStyle`` IR operation to a DOCX document.

    Looks up the named style and applies all properties from the IR.
    Raises ``ValueError`` if the style does not exist.

    Args:
        doc: The python-docx ``Document`` to mutate.
        ir: The ``EditStyle`` instruction.

    Raises:
        ValueError: If *ir.style_name* does not exist in *doc.styles*.
    """
    try:
        style = doc.styles[ir.style_name]
    except KeyError:
        raise ValueError(
            f"Style '{ir.style_name}' not found in document styles."
        ) from None

    # Convert dict properties to StyleProperties if needed
    if isinstance(ir.properties, dict):
        properties = StyleProperties(**ir.properties)
    else:
        properties = ir.properties

    _apply_style_properties(style, properties)


_TABLE_ALIGNMENT_MAP: dict[str, WD_TABLE_ALIGNMENT] = {
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}

_EDGES = ("top", "left", "bottom", "right", "insideH", "insideV")


def apply_table_format(doc: Any, ir: ApplyTableFormat) -> None:
    """Apply ``ApplyTableFormat`` IR operation to a DOCX document.

    For each table index, sets table alignment and injects ``w:tblBorders``
    XML into the table's ``w:tblPr``. Any existing ``w:tblBorders`` is
    removed first to avoid duplication.

    Args:
        doc: The python-docx ``Document`` to mutate.
        ir: The ``ApplyTableFormat`` instruction.
    """
    for table_idx in ir.table_indices:
        table = doc.tables[table_idx]

        tbl_pr = _ensure_child(table._tbl, "w:tblPr")

        if ir.format.alignment is not None:
            table.alignment = _TABLE_ALIGNMENT_MAP.get(
                ir.format.alignment, WD_TABLE_ALIGNMENT.CENTER
            )

        if ir.format.borders is not None:

            existing = tbl_pr.find(qn("w:tblBorders"))
            if existing is not None:
                tbl_pr.remove(existing)

            tbl_borders = OxmlElement("w:tblBorders")
            borders = ir.format.borders

            for edge in _EDGES:
                edge_el = OxmlElement(f"w:{edge}")
                edge_el.set(qn("w:val"), borders.val)
                edge_el.set(qn("w:sz"), str(borders.sz))
                edge_el.set(qn("w:space"), str(borders.space))
                edge_el.set(qn("w:color"), borders.color)
                tbl_borders.append(edge_el)

            tbl_pr.append(tbl_borders)
